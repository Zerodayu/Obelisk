"""Owned local services for the legacy live scripts; no production configuration."""
import asyncio
from contextlib import redirect_stdout
import importlib.util
import io
import json
import os
from pathlib import Path
import socket
import subprocess
import sys
import tempfile
import time
import uuid

ROOT = Path(__file__).resolve().parents[2]
HERE = Path(__file__).resolve().parent
FIXTURES = tuple(ROOT / "classrecord_templates" / f"JMCFI_Class_Record_Template_AUN-OBE({kind}).xlsx"
                 for kind in ("worst", "mixed", "ideal"))
LIVE_CASES = ("local service preflight", "upload E2E worst/mixed/ideal", "stress health/analytics/upload",
              "all accepted Redis jobs completed", "invalid workbook worker rejection")
OFFLINE_CASES = ("offline unittest", "standalone validation", "standalone CQI (mock Gemini)", "mocked Redis mode safety")


class Blocked(RuntimeError):
    pass


def _preflight():
    if any(os.getenv(name) for name in ("HTTP_PROXY", "HTTPS_PROXY", "ALL_PROXY", "http_proxy", "https_proxy", "all_proxy")):
        raise Blocked("Proxy environment is unsupported for isolated localhost tests")
    for package in ("uvicorn", "httpx", "requests", "redis", "openpyxl"):
        if importlib.util.find_spec(package) is None:
            raise Blocked(f"Missing dependency: {package}; dependencies were not installed")
    for path in FIXTURES:
        if not path.is_file():
            raise Blocked(f"Required fixture missing: {path.name}")
    from testing_modules.old_test._redis_ping import _verify_local_container, _docker
    _verify_local_container()  # Read-only ownership check; never PING or use shared Redis.
    context = _docker("context", "show")
    image = _docker("--context", context, "inspect", "obelisk-redis", "--format", "{{.Image}}")
    if not image.startswith("sha256:"):
        raise Blocked("Cached Redis image identity could not be verified")
    return context, image


def _server(scratch, port, redis_port, token, site_packages):
    """Fresh interpreter: real production routers/workers, isolated test settings."""
    sys.dont_write_bytecode = True
    sys.path[:0] = [str(ROOT), site_packages]
    scratch = Path(scratch).resolve()
    # Starlette may spool large multipart bodies before the upload route runs.
    tempfile.tempdir = str(scratch)
    for key in list(os.environ):
        if key.upper().startswith("OBELISK_"):
            del os.environ[key]
    os.environ.update(OBELISK_ENV="legacy_live_" + token, OBELISK_REDIS_HOST="127.0.0.1",
                      OBELISK_REDIS_PORT=str(redis_port), OBELISK_UPLOAD_FOLDER=str(scratch / "uploads"))
    # Only this child process is configured. Root env files are not read.
    import app.core.env as env
    env.load_env = lambda: None
    import app.analytics.cqi_recommender as cqi
    cqi.IS_DEBUG_MODE = True
    def denied_provider(*args, **kwargs):
        raise RuntimeError("Real Gemini clients are forbidden by --all-live")
    cqi.genai.Client = denied_provider
    import app.main as main
    import uvicorn
    from fastapi import Header, HTTPException
    from app.services import job_queue

    def guard(event, args):
        if event == "socket.connect":
            caller = sys._getframe(1)
            if caller.f_code.co_name == "_fallback_socketpair" and Path(caller.f_code.co_filename).name == "socket.py":
                return
            if args[1] != ("127.0.0.1", redis_port):
                raise RuntimeError("Live test server blocked non-owned outbound connection")
        elif event == "socket.getaddrinfo":
            host = args[0].decode() if isinstance(args[0], bytes) else args[0]
            if host != "127.0.0.1":
                raise RuntimeError("Live test server blocked DNS")
        elif event in {"subprocess.Popen", "os.system"}:
            raise RuntimeError("Live test server cannot spawn descendants")
        elif event == "open":
            path, mode, flags = args
            writing = isinstance(mode, str) and any(c in mode for c in "wax+")
            writing |= isinstance(flags, int) and bool(flags & (os.O_WRONLY | os.O_RDWR | os.O_CREAT | os.O_TRUNC))
            if writing and isinstance(path, (str, bytes, os.PathLike)) and not Path(os.fsdecode(path)).resolve().is_relative_to(scratch):
                raise RuntimeError("Live test server blocked persistent write")
        elif event in {"os.mkdir", "os.remove", "os.rmdir", "os.rename"}:
            paths = args[:2] if event == "os.rename" else args[:1]
            if any(not Path(os.fsdecode(p)).resolve().is_relative_to(scratch) for p in paths):
                raise RuntimeError("Live test server blocked filesystem mutation")
    sys.addaudithook(guard)
    server = uvicorn.Server(uvicorn.Config(main.app, host="127.0.0.1", port=port,
                                         log_level="warning", access_log=False))
    def owned(value):
        if value != token:
            raise HTTPException(403, "Owned test server token required")
    @main.app.get("/__legacy_probe")
    async def probe(x_legacy_token: str = Header(default="")):
        owned(x_legacy_token)
        return dict(pid=os.getpid(), token=token, redis_port=redis_port, mock_gemini=cqi.IS_DEBUG_MODE,
                    workers=len(main._worker_tasks), redis_ready=bool(job_queue.redis_client),
                    multipart_temp_root=tempfile.gettempdir())
    @main.app.post("/__legacy_stop")
    async def stop(x_legacy_token: str = Header(default="")):
        owned(x_legacy_token)
        server.should_exit = True
        return {"stopping": True}
    server.run()


def _offline(rows):
    result = subprocess.run([sys.executable, "-B", str(HERE / "run_all.py")], cwd=ROOT,
                            capture_output=True, text=True, encoding="utf-8", errors="replace", timeout=120)
    for name, marker in (("offline unittest", "UNITTEST RESULT: 50 run, 0 failures, 0 errors, 0 skipped"),
                         ("standalone validation", "PASS: test_validate (exit 0)"),
                         ("standalone CQI (mock Gemini)", "PASS: test_ai_module (exit 0)")):
        rows[name] = ("PASS" if marker in result.stdout else "FAIL", marker)
    if result.returncode:
        print(result.stdout[-4000:] + result.stderr[-2000:])
    safety = subprocess.run([sys.executable, "-B", "-m", "unittest", "testing_modules.old_test.test_redis_ping",
                             "testing_modules.old_test.test_all_live_safety"],
                            cwd=ROOT, capture_output=True, text=True, timeout=30)
    rows["mocked Redis mode safety"] = ("PASS" if safety.returncode == 0 else "FAIL", "12 Redis + 6 all-live safety tests, separately from default 50")


def run_all_live():
    print("WARNING: --all-live creates disposable local Redis/server/uploads and bounded HTTP load.", flush=True)
    print("Gemini is MOCKED and real clients denied; no paid API calls, shared queue or deployed service.", flush=True)
    print("Stress profile: concurrency 2; 4 health + 20 analytics + 15 uploads (legacy minimums).", flush=True)
    rows = {name: ("BLOCKED", "Not executed: prerequisites or earlier setup incomplete") for name in OFFLINE_CASES + LIVE_CASES}
    rows["real Gemini/provider evaluation"] = ("SKIP", "Outside --all-live scope: mock Gemini is mandatory")
    container = None
    process = None
    client = None
    log = None
    context = None
    token = uuid.uuid4().hex
    forced = False
    scratch_manager = tempfile.TemporaryDirectory(prefix="obelisk-all-live-")
    scratch = Path(scratch_manager.name)
    try:
        _offline(rows)
        context, image = _preflight()
        print("PASS: offline phase finished; starting owned local Redis/server (Gemini MOCKED).", flush=True)
        from testing_modules.old_test._redis_ping import _docker
        import httpx
        import redis
        container = "obelisk-legacy-" + token
        _docker("--context", context, "run", "--detach", "--pull=never", "--name", container,
                "--label", "obelisk.legacy.owner=" + token, "--publish", "127.0.0.1::6379",
                image, "redis-server", "--save", "", "--appendonly", "no")
        metadata = json.loads(_docker("--context", context, "inspect", container, "--format",
                                     '{"ports":{{json .NetworkSettings.Ports}},"labels":{{json .Config.Labels}}}'))
        if metadata["labels"].get("obelisk.legacy.owner") != token:
            raise Blocked("Disposable Redis ownership mismatch")
        bindings = metadata["ports"]["6379/tcp"]
        if len(bindings) != 1 or bindings[0]["HostIp"] != "127.0.0.1":
            raise Blocked("Disposable Redis is not published only on loopback")
        redis_port = int(bindings[0]["HostPort"])
        if redis_port == 6379:
            raise Blocked("Refusing shared Redis port")
        with redis.Redis(host="127.0.0.1", port=redis_port, socket_timeout=1, socket_connect_timeout=1) as db:
            if not db.ping() or db.dbsize() != 0:
                raise Blocked("Disposable Redis is not ready/empty")
        with socket.socket() as reservation:
            reservation.bind(("127.0.0.1", 0))
            port = reservation.getsockname()[1]
        url = f"http://127.0.0.1:{port}"
        log = (scratch / "server.log").open("w", encoding="utf-8")
        # Launch the base interpreter directly to avoid Windows venv redirector descendants.
        executable = getattr(sys, "_base_executable", sys.executable)
        site_packages = str(Path(sys.prefix) / ("Lib/site-packages" if os.name == "nt" else f"lib/python{sys.version_info.major}.{sys.version_info.minor}/site-packages"))
        process = subprocess.Popen([executable, "-B", str(Path(__file__).resolve()), "--serve", str(scratch),
                                    str(port), str(redis_port), token, site_packages], cwd=ROOT,
                                   stdout=log, stderr=log,
                                   creationflags=subprocess.CREATE_NO_WINDOW if os.name == "nt" else 0)
        client = httpx.Client(base_url=url, timeout=5, trust_env=False)
        deadline = time.monotonic() + 20
        while True:
            if process.poll() is not None:
                raise Blocked("Owned test server exited during startup")
            try:
                proof = client.get("/__legacy_probe", headers={"X-Legacy-Token": token}).json()
                if (proof.get("pid") == process.pid and proof.get("token") == token
                        and proof.get("redis_port") == redis_port and proof.get("mock_gemini") is True
                        and proof.get("redis_ready") is True and proof.get("workers") == 4
                        and proof.get("multipart_temp_root") == str(scratch.resolve())):
                    break
            except (httpx.HTTPError, ValueError):
                pass
            if time.monotonic() >= deadline:
                raise Blocked("Owned server readiness/worker proof timed out")
            time.sleep(0.2)
        rows["local service preflight"] = ("PASS", "owned PID, 4 real workers, fresh Redis, mock Gemini")
        print("PASS: owned server and 4 real Redis workers ready; running original E2E scripts.", flush=True)
        from testing_modules.old_test import test_upload_e2e as upload, test_stress_async as stress
        upload.BASE_URL = url  # Original script has no --url; assertions and --live remain intact.
        with redirect_stdout(io.StringIO()) as output:
            status = upload.main(["--live"])
        rows["upload E2E worst/mixed/ideal"] = ("PASS" if status == 0 else "FAIL", "original script --live; actual HTTP/Redis/ETL, mock CQI")
        for fixture in FIXTURES:
            passed = f"{fixture.name}: PASSED" in output.getvalue()
            rows["upload " + fixture.name] = ("PASS" if passed else "FAIL", "original E2E assertions; Gemini mocked")
        if status:
            print(output.getvalue()[-4000:])
        with redirect_stdout(io.StringIO()) as output:
            status = stress.main(["--live", "--url", url, "--mode", "all", "--concurrency", "2", "--requests", "4"])
        rows["stress health/analytics/upload"] = ("PASS" if status == 0 else "FAIL", "4/20/15 requests; HTTP success only, completion checked separately")
        if status:
            print(output.getvalue()[-4000:])
        deadline = time.monotonic() + 45
        while True:
            response = client.get("/jobs/")
            response.raise_for_status()
            jobs = response.json()
            if all(j["status"] in {"completed", "failed"} for j in jobs):
                break
            if time.monotonic() >= deadline:
                raise RuntimeError("Accepted jobs did not drain within 45 seconds")
            time.sleep(0.2)
        from openpyxl import load_workbook
        expected_counts = {}
        for path in FIXTURES:
            book = load_workbook(path, read_only=False, data_only=True)
            try:
                sheet = book["Direct CLO"]
                clos = sum(bool(sheet.cell(3, col).value) for col in range(3, sheet.max_column + 1, 7))
                roster = 0
                for row in sheet.iter_rows(min_row=5, max_col=2, values_only=True):
                    if not any(row) or str(row[1] or "").startswith("CLASS AVERAGE"):
                        break
                    roster += 1
                expected_counts[path.name] = roster * clos
            finally:
                book.close()
        namespace = "obelisk:legacy_live_" + token + ":"
        verified = len(jobs) == 18
        with redis.Redis(host="127.0.0.1", port=redis_port, decode_responses=True, socket_timeout=1) as db:
            verified &= db.llen(namespace + "job_queue") == 0
            for job in jobs:
                loaded = (job.get("result") or {}).get("loaded") or {}
                persisted = json.loads(db.get(namespace + "job:" + job["job_id"]) or "null")
                verified &= (job["status"] == "completed" and persisted == job
                             and len(loaded.get("attainments", [])) == expected_counts.get(job["payload"]["filename"], -1)
                             and loaded.get("received_records") == len(loaded.get("attainments", []))
                             and Path(job["payload"]["file_path"]).resolve().is_relative_to(scratch.resolve()))
        rows["all accepted Redis jobs completed"] = ("PASS" if verified else "FAIL", f"{len(jobs)}/18 jobs; persisted results, expected record counts, queue drained")
        invalid = client.post("/upload", files={"file": ("invalid.xlsx", b"not an Excel workbook")})
        invalid.raise_for_status()
        job_id = invalid.json()["job_id"]
        deadline = time.monotonic() + 15
        while True:
            job = client.get("/jobs/" + job_id).json()
            if job["status"] in {"completed", "failed"}:
                break
            if time.monotonic() >= deadline:
                raise RuntimeError("Invalid workbook did not reach terminal state")
            time.sleep(0.2)
        valid_rejection = invalid.status_code == 202 and job["status"] == "failed" and job["error"]["error_type"] == "InvalidWorkbook" and job["result"] is None
        rows["invalid workbook worker rejection"] = ("PASS" if valid_rejection else "FAIL", "real upload accepted then structured worker failure; no result")
    except Exception as exc:
        blocked = isinstance(exc, (Blocked, subprocess.TimeoutExpired)) or type(exc).__name__ == "RedisCheckBlocked"
        rows["harness"] = ("BLOCKED" if blocked else "FAIL", str(exc))
        if blocked:
            for name in OFFLINE_CASES + LIVE_CASES:
                if rows[name][0] == "BLOCKED":
                    rows[name] = ("BLOCKED", str(exc))
    finally:
        try:
            if process is not None and process.poll() is None:
                try:
                    client.post("/__legacy_stop", headers={"X-Legacy-Token": token})
                    process.wait(timeout=10)
                except Exception:
                    forced = True
                    process.kill()
                    process.wait(timeout=5)
            if process is not None and process.returncode != 0:
                forced = True
            if client is not None:
                client.close()
            if log is not None:
                log.close()
            if container is not None:
                from testing_modules.old_test._redis_ping import _docker
                owner = _docker("--context", context, "inspect", container, "--format", '{{index .Config.Labels "obelisk.legacy.owner"}}')
                if owner != token:
                    raise Blocked("Refusing cleanup of a container with mismatched ownership")
                _docker("--context", context, "rm", "--force", container)
                remaining = _docker("--context", context, "ps", "--all", "--quiet", "--filter", "label=obelisk.legacy.owner=" + token)
                if remaining:
                    raise Blocked("Owned container cleanup incomplete")
            scratch_manager.cleanup()
            rows["owned service/file cleanup"] = ("BLOCKED" if forced else "PASS", "server stopped; owned container/temp files removed" if not forced else "forced process stop; graceful shutdown not verified")
        except Exception as exc:
            rows["owned service/file cleanup"] = ("BLOCKED", str(exc))
    print("\nALL-LIVE RESULTS (real HTTP/Redis/workers; Gemini MOCKED):")
    for name, (status, detail) in rows.items():
        print(f"{status}: {name}: {detail}")
    return 1 if any(status == "FAIL" for status, _ in rows.values()) else 2 if any(status == "BLOCKED" for status, _ in rows.values()) else 0


if __name__ == "__main__":
    if len(sys.argv) == 7 and sys.argv[1] == "--serve":
        _server(sys.argv[2], int(sys.argv[3]), int(sys.argv[4]), sys.argv[5], sys.argv[6])
    else:
        raise SystemExit("Use run_all.py --all-live; direct harness invocation is unsupported")
