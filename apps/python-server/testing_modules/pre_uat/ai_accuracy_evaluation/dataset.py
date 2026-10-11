"""Independent workbook arithmetic; no application, service or provider imports."""
from collections import Counter
from decimal import Decimal, ROUND_HALF_EVEN
import hashlib
import math
from openpyxl import load_workbook

DATASET = 'classrecord_templates/JMCFI_Class_Record_Template_AUN-OBE(mixed).xlsx'
KNOWN_SYNTHETIC_SHA256 = '867873f6fd2dd016e39ffc8a7371dc5e5eff618e74c894eb90865690cb1c8e73'


def number(value):
    if isinstance(value, bool) or not isinstance(value, (int, float)) or not math.isfinite(value):
        raise ValueError('Expected a finite numeric workbook cell')
    return Decimal(str(value))


def inspect(root):
    path = root / DATASET
    sha = hashlib.sha256(path.read_bytes()).hexdigest()
    if sha != KNOWN_SYNTHETIC_SHA256:
        raise ValueError('Workbook differs from the inspected synthetic dataset; refusing live disclosure')
    book = load_workbook(path, read_only=True, data_only=True)
    formulas = load_workbook(path, read_only=True, data_only=False)
    try:
        grids = {name: list(book[name].values) for name in ('Direct CLO', 'Indirect CLO', 'RECORD', 'F12', 'SETUP')}
        direct, indirect, source, survey = [grids[n] for n in ('Direct CLO', 'Indirect CLO', 'RECORD', 'F12')]
        assert direct[0][0] == 'DIRECT CLO — PER-STUDENT RAW SCORES'
        assert indirect[0][0] == 'INDIRECT CLO — COURSE EXIT SURVEY RESPONSES'
        blocks = [(column, str(direct[2][column]).strip()) for column in range(2, len(direct[2]), 7) if direct[2][column]]
        assert [clo for _, clo in blocks] == ['CLO1', 'CLO2', 'CLO3', 'CLO4', 'CLO5']
        roster = []
        for row, cells in enumerate(direct[4:], 5):
            if not any(cells[:2]) or str(cells[1] or '').upper().startswith('CLASS AVERAGE'):
                break
            roster.append((row, cells))
        assert len(roster) == 30 and len({cells[0] for _, cells in roster}) == 30
        indirect_map = {row[0]: row for row in indirect[4:34]}
        assert set(indirect_map) == {cells[0] for _, cells in roster}
        cases, identity_map, coverage = [], {}, {}
        for ordinal, (row, cells) in enumerate(roster, 1):
            student = f'S{ordinal:03}'
            identity_map[(cells[0], cells[1])] = student
            assert cells[:2] == source[row + 8][1:3]
            assert cells[0] == survey[row - 1][1]
            for column, clo in blocks:
                raw = cells[column:column + 6]
                rating = indirect_map[cells[0]][2 + 2 * (int(clo[-1]) - 1)]
                assert rating == survey[row - 1][3 + int(clo[-1]) - 1]
                assert 1 <= number(rating) <= 5
                periods = []
                for k, period in enumerate(('Prelim', 'Midterm', 'Final')):
                    cols = [j for j in range(4, 148) if j < len(source[11]) and source[11][j] == clo
                            and j < len(source[10]) and source[10][j] == period]
                    score = sum((number(source[row + 8][j]) for j in cols), Decimal(0))
                    maximum = sum((number(source[12][j]) for j in cols), Decimal(0))
                    assert score == number(raw[2 * k]) and maximum == number(raw[2 * k + 1])
                    assert 0 <= score <= maximum
                    coverage.setdefault(clo, {})[period] = {'assessment_count': len(cols), 'max_score': float(maximum)}
                    periods.append({'period': period, 'score': float(score), 'max_score': float(maximum)})
                total_score = sum((number(p['score']) for p in periods), Decimal(0))
                total_max = sum((number(p['max_score']) for p in periods), Decimal(0))
                assert total_max > 0
                direct_pct = total_score / total_max * 100
                indirect_pct = number(rating) / 5 * 100
                composite = (Decimal('.7') * direct_pct + Decimal('.3') * indirect_pct).quantize(Decimal('.01'), rounding=ROUND_HALF_EVEN)
                cases.append(dict(case_id=f'{student}:{clo}', student=student, clo_code=clo, workbook_row=row,
                                  periods=periods, rating=float(rating), direct_pct=float(direct_pct),
                                  indirect_pct=float(indirect_pct), composite_pct=float(composite),
                                  at_risk=composite < 70, complete_nonnull=all(x is not None for x in raw[::2])))
        clos = []
        for _, code in blocks:
            selected = [c for c in cases if c['clo_code'] == code]
            mean = sum((Decimal(str(c['composite_pct'])) for c in selected), Decimal(0)) / len(selected)
            clos.append(dict(clo_code=code, eligible_records=len(selected), mean_composite_pct=float(mean),
                             at_risk=mean < 70, below_threshold_records=sum(c['at_risk'] for c in selected)))
        assert len(cases) == 150
        manifest = dict(path=DATASET, sha256=sha, synthetic=True, students=30, clos=5,
                        valid_student_clo_cases=len(cases), missing_or_invalid_cases=0,
                        sheets=book.sheetnames, assessment_coverage=coverage,
                        independent_source_validation='RECORD raw assessments and F12 ratings match Direct/Indirect caches',
                        rules={'direct': 'sum raw / sum maxima, converted to percent', 'indirect': 'rating / 5 * 100',
                               'composite': '0.70 * direct_pct + 0.30 * indirect_pct; two decimals, half-even',
                               'student_positive': 'rounded composite < 70; exactly 70 passes',
                               'clo_positive': 'mean of rounded student composite values < 70; PROVISIONAL',
                               'completeness': 'non-null period scores; zero assessment coverage disclosed separately'},
                        limitations=['CLO4 and CLO5 have Final assessments only; zero Prelim/Midterm maxima.',
                                     'Workbook formula caches verified against raw sources; Excel itself was not run.',
                                     'One synthetic course; no population-level generalization.'])
        return cases, clos, manifest, identity_map
    finally:
        book.close()
        formulas.close()
