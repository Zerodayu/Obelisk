"""ST-AI-001 only. Default: credential-safe preflight, zero generation calls.

--execute calls the unchanged production recommender only if its existing
configuration/SDK can satisfy the single-request bound. No provider replacement,
mock, retry override or network blocker is installed.
"""
import argparse
import ast
import asyncio
import contextlib
from datetime import datetime, timezone
from decimal import Decimal, ROUND_HALF_EVEN
import hashlib
from importlib.metadata import version
import io
import json
import logging
from pathlib import Path
import platform
import re
import sys
import tempfile
import time
from types import ModuleType
import unittest
from urllib.parse import urlparse

ROOT = Path(__file__).resolve().parents[2]
sys.path.insert(0, str(ROOT))
sys.dont_write_bytecode = True


def fixture_helpers(cqi):
    """Reuse only fixture/oracle definitions, without offline bootstrap or mocks.

    support.py normally requires Phase 2A's network-denying bootstrap. Extract
    its unchanged named definitions instead; never execute its module entry or
    mocked_recommendation. Its legacy helper reads the unchanged historical
    workbook builder without running historical test methods.
    """
    from openpyxl import load_workbook
    from app.etl import etl_const
    from app.etl.extract.extractor import ExcelExtractor
    from app.etl.transform.transformer import SimpleTransformer

    source = ROOT / "testing_modules/pre_uat/support.py"
    tree = ast.parse(source.read_text(encoding="utf-8-sig"))
    names = {"legacy", "SyntheticWorkbook", "oracle", "gap_oracle", "check_numeric_claims"}
    tree.body = [node for node in tree.body if
                 isinstance(node, (ast.FunctionDef, ast.ClassDef)) and node.name in names
                 or isinstance(node, ast.Assign) and any(
                     isinstance(target, ast.Name) and target.id == "CLAIM"
                     for target in node.targets)]
    namespace = dict(ROOT=ROOT, LEGACY={}, ModuleType=ModuleType, ast=ast,
                     tempfile=tempfile, Path=Path, load_workbook=load_workbook,
                     C=etl_const, ExcelExtractor=ExcelExtractor,
                     SimpleTransformer=SimpleTransformer, CQI=cqi,
                     Decimal=Decimal, ROUND_HALF_EVEN=ROUND_HALF_EVEN, re=re)
    exec(compile(tree, str(source), "exec"), namespace)
    return namespace


def load_configuration():
    # Production provider logs contain key previews. Discard them rather than
    # displaying partial credentials; never save the captured logging buffer.
    logging.disable(logging.CRITICAL + 1)
    discarded = io.StringIO()
    with contextlib.redirect_stdout(discarded), contextlib.redirect_stderr(discarded):
        from app.analytics import cqi_recommender as cqi
    return cqi


def preflight(cqi):
    keys = cqi.settings.llm_api_keys_list
    checks = dict(debug_mode=cqi.IS_DEBUG_MODE, configured_key_count=len(keys),
                  configured_model=cqi.settings.LLM_MODEL,
                  google_genai_version=version("google-genai"),
                  configured_key_values_recorded=False)
    blockers = []
    if cqi.IS_DEBUG_MODE:
        blockers.append("Production IS_DEBUG_MODE is enabled; configuration will not be changed.")
    if len(keys) != 1:
        blockers.append("Exactly one configured key is required: production automatically tries each key after failures.")
    if not keys:
        return checks, blockers
    # Construct the exact production SDK client; initialization sends no
    # generation request. Inspect its actual defaults instead of guessing from
    # HttpRetryOptions documentation. Nothing on this client is modified.
    client = cqi.genai.Client(api_key=keys[0])
    try:
        api = client._api_client
        attempts = api._async_retry.stop.max_attempt_number
        options = api._http_options
        uses_aiohttp = api._use_aiohttp()
        checks.update(sdk_max_attempts=attempts,
                      sdk_retry_options_supplied=options.retry_options is not None,
                      async_transport="aiohttp" if uses_aiohttp else "httpx",
                      provider_host=urlparse(options.base_url).hostname,
                      vertexai=bool(api.vertexai))
        if attempts != 1:
            blockers.append("Existing SDK retry configuration permits multiple requests.")
        if uses_aiohttp:
            blockers.append("Installed SDK aiohttp path retries connection failures internally; one request cannot be guaranteed unchanged.")
        if api.vertexai or checks["provider_host"] != "generativelanguage.googleapis.com":
            blockers.append("Existing SDK target is not the direct Gemini API endpoint; no endpoint override will be applied.")
    finally:
        client.close()
    return checks, blockers


async def execute(cqi, record):
    helpers = fixture_helpers(cqi)
    with helpers["SyntheticWorkbook"](students=4, clos=3, scenario="worst") as fixture:
        extracted, output = await fixture.etl()
        assert len(output) == 12
        assert len({(a.student_id, a.clo_code) for a in output}) == 12
        assert len({a.student_id for a in output}) == 4
        assert {a.clo_code for a in output} == {"CLO1", "CLO2", "CLO3"}
        # Independent expected arithmetic, not the production calculation helper.
        expected_composite = float(Decimal("0.7") * 0 + Decimal("0.3") * 20)
        for raw, attainment in zip(extracted[1], output, strict=True):
            expected = helpers["oracle"](raw)
            assert expected["direct"] == attainment.direct_clo_attainment_pct == 0.0
            assert expected["indirect"] == attainment.indirect_clo_attainment_pct == 20.0
            assert expected["composite"] == attainment.composite_clo_attainment_pct == expected_composite
            assert attainment.met_threshold is False
        before = [a.model_dump() for a in output]
        gap = cqi.identify_gaps(extracted[0], cqi.anonymize_students(output))
        assert all(g["num_students_below_threshold"] == g["total_students"] == 4
                   and g["threshold"] == 70.0 for g in gap["gaps"])
        record["synthetic_input"] = dict(students=4, clos=3, records=12,
                                        direct_pct=0, indirect_pct=20, composite_pct=6,
                                        proficiency_threshold_pct=70, below_threshold_records=12)
        record["actual_prompt"] = cqi.build_prompt(extracted[0], gap)
        record["input_records"] = [a.model_dump() for a in cqi.anonymize_students(output)]
        record["fixture_sha256"] = hashlib.sha256(fixture.path.read_bytes()).hexdigest()
        # Passive Python profiling observes genuine SDK frames/return values.
        # It neither intercepts requests nor changes any provider function.
        observations = {"provider_invoked": False, "sdk_generation_invoked": False,
                        "sdk_http_attempt_frames": set()}
        def observe(frame, event, value):
            name = frame.f_code.co_name
            filename = frame.f_code.co_filename.replace("\\", "/")
            if event == "call" and name == "call_llm_api" and filename.endswith("/app/analytics/cqi_recommender.py"):
                observations["provider_invoked"] = True
                record["actual_prompt"] = frame.f_locals["prompt"]
            if event == "call" and name == "generate_content" and filename.endswith("/google/genai/models.py"):
                observations["sdk_generation_invoked"] = True
                record["requested_model"] = frame.f_locals.get("model")
                assert frame.f_locals.get("contents") == record["actual_prompt"]
            if event == "call" and name == "_async_request_once" and filename.endswith("/google/genai/_api_client.py"):
                observations["sdk_http_attempt_frames"].add(id(frame))
            if event == "return" and name == "generate_content" and isinstance(value, cqi.types.GenerateContentResponse):
                record["actual_generated_response"] = value.text
                record["response_model_version"] = value.model_version
                record["provider_response_id"] = value.response_id
        previous = sys.getprofile()
        sys.setprofile(observe)
        record["production_recommender_invoked"] = True
        try:
            result = await asyncio.wait_for(
                cqi.generate_cqi_recommendation(extracted[0], output), timeout=60)
        finally:
            sys.setprofile(previous)
            record.update(provider_invoked=observations["provider_invoked"],
                          sdk_generation_invoked=observations["sdk_generation_invoked"],
                          observed_sdk_http_attempts=len(observations["sdk_http_attempt_frames"]))
        assert before == [a.model_dump() for a in output]
        record["inputs_unchanged"] = True
        record["production_result"] = result
        if result["status"] != "ok" or not record.get("actual_generated_response"):
            record.update(status="BLOCKED", reason="Provider did not return a successful generated course recommendation; see sanitized result.")
            if record.get("actual_generated_response"):
                record.update(status="FAIL", reason="Gemini generated content but production rejected the recommendation schema.")
            return
        assert record["provider_invoked"] and record["sdk_generation_invoked"]
        assert record["observed_sdk_http_attempts"] == 1
        assert cqi.validate_course_recommendation(result["recommendation"])
        record["schema_validation"] = "PASS — actual production course schema"
        data = json.loads(result["recommendation"])
        grounding = helpers["check_numeric_claims"](data, helpers["gap_oracle"](extracted[1]))
        record["numerical_grounding"] = grounding
        assert not grounding["mismatches"]
        record["responsible_roles"] = [r["recommendedFor"] for r in data["recommendations"]]
        record["manual_review"] = [
            "Confirm recommendations address the three documented CLO gaps with useful improvement actions.",
            "Confirm wording remains advisory and does not recommend altering/inflating grades.",
            "Review every numerical claim outside the documented exact-sentence grammar, including numbers written as words.",
            "Confirm academic responsibility and educational usefulness; role enum validity alone is insufficient."]
        record.update(status="PASS", pass_scope="Live generation, schema, exact-grammar numerical checks and unchanged inputs only",
                      requirement_satisfied="Generation verified; full semantic/advisory/grounding acceptance requires human review.")


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--execute", action="store_true", help="Explicitly request the one live execution after preflight")
    args = parser.parse_args()
    started = time.perf_counter()
    stamp = datetime.now(timezone.utc)
    destination = ROOT / "testing_modules/pre_uat/evidence" / (stamp.strftime("%Y%m%dT%H%M%S_%fZ") + "_ST-AI-001_live")
    destination.mkdir(exist_ok=False)
    record = dict(test_id="ST-AI-001", test_type="Live Gemini AI generation",
                  timestamp=stamp.isoformat(), status="BLOCKED", live_execution_requested=args.execute,
                  production_recommender_invoked=False, provider_invoked=False,
                  sdk_generation_invoked=False, observed_sdk_http_attempts=0,
                  expected="Generate valid advisory CQI recommendations for 12 records at 6% composite attainment, all below 70%.",
                  environment=dict(python=platform.python_version(), platform=platform.platform(),
                                   database="not imported or accessed", redis="not imported or accessed",
                                   workers="not started", provider="production Gemini SDK; no mocks or replacement provider"))
    cqi = None
    discarded = io.StringIO()
    try:
        cqi = load_configuration()
        with contextlib.redirect_stdout(discarded), contextlib.redirect_stderr(discarded):
            checks, blockers = preflight(cqi)
            record["preflight"] = checks
            if blockers:
                record["reason"] = " ".join(blockers)
            elif not args.execute:
                record.update(status="READY", reason="Preflight passed. No Gemini generation requested in preflight mode.")
            else:
                asyncio.run(execute(cqi, record))
    except Exception as exc:
        # Never persist str(exc) or traceback: settings/provider exceptions can
        # contain credentials. HTTP codes/classes suffice for sanitized diagnosis.
        record.update(status="FAIL" if isinstance(exc, AssertionError) else "BLOCKED",
                      error_type=type(exc).__name__,
                      reason="Numerical/contract assertion failed." if isinstance(exc, AssertionError)
                      else "Configuration, SDK preflight or provider execution could not complete.")
        code = getattr(exc, "code", None)
        if isinstance(code, int):
            record["http_error_code"] = code
    record["duration_seconds"] = round(time.perf_counter() - started, 6)
    # Remove full keys and key previews from all serialized provider errors.
    serialized = json.dumps(record, indent=2, ensure_ascii=False)
    if cqi is not None:
        for key in cqi.settings.llm_api_keys_list:
            serialized = serialized.replace(key, "[REDACTED]")
            if len(key) > 8:
                serialized = serialized.replace(f"{key[:4]}...{key[-4:]}", "[REDACTED]")
    (destination / "ST-AI-001_result.txt").write_text(serialized, encoding="utf-8")
    sources = [ROOT / "app/analytics/cqi_recommender.py", ROOT / "app/core/config.py",
               ROOT / "testing_modules/pre_uat/support.py", ROOT / "testing_modules/old_test/test_etl_v2.py",
               Path(__file__)]
    (destination / "SOURCE_MANIFEST.json").write_text(json.dumps(
        {p.relative_to(ROOT).as_posix(): hashlib.sha256(p.read_bytes()).hexdigest() for p in sources}, indent=2), encoding="utf-8")
    safe = json.loads(serialized)
    print("ST-AI-001 | LIVE GEMINI | " + safe["status"])
    print("Production recommender invoked: " + str(safe["production_recommender_invoked"]))
    print("Observed SDK HTTP attempts: " + str(safe["observed_sdk_http_attempts"]))
    print("Actual result: " + safe.get("reason", safe.get("requirement_satisfied", "See evidence")))
    if safe.get("actual_generated_response"):
        print("ACTUAL GEMINI-GENERATED RESPONSE:\n" + safe["actual_generated_response"])
    print("Evidence: " + destination.relative_to(ROOT).as_posix())
    return 0 if safe["status"] in {"READY", "PASS"} else 2


if __name__ == "__main__":
    raise SystemExit(main())
