"""Synthetic fixtures, independent oracle, and evidence helpers for Phase 2A."""
import ast
import asyncio
from collections import Counter
from decimal import Decimal, ROUND_HALF_EVEN
import inspect
from functools import wraps
import json
from pathlib import Path
import re
import subprocess
import tempfile
from types import ModuleType, SimpleNamespace
import unittest
from unittest.mock import AsyncMock, patch

from testing_modules.pre_uat.isolation import ROOT, STATE

if not STATE["verified"]:
    raise RuntimeError("Run through pre_uat/run_phase2a.py; offline isolation is required")

from openpyxl import load_workbook
from app.etl import etl_const as C
from app.etl.extract.extractor import ExcelExtractor
from app.etl.transform.transformer import SimpleTransformer
from app.analytics import cqi_recommender as CQI

REVISION = "1a8fb6dd41d1ef9c10fe5ab0f03e24af54a3ef93"
LEGACY = {}
REUSED = []


def legacy(name):
    """Read restored originals; execute only imports and class definitions in memory.

    Never import a live/standalone test script or change its on-disk original.
    """
    if name not in LEGACY:
        if name not in {"test_etl_v2.py", "test_composite_attainment.py"}:
            raise ValueError("Historical source not allowlisted")
        path = ROOT / "testing_modules" / "old_test" / name
        source = path.read_text(encoding="utf-8-sig")
        tree = ast.parse(source)
        tree.body = [node for node in tree.body if isinstance(node, (ast.Import, ast.ImportFrom, ast.ClassDef))]
        module = ModuleType("pre_uat_historical_" + name[:-3])
        module.__file__ = str(path)
        exec(compile(tree, str(path), "exec"), module.__dict__)
        LEGACY[name] = module
    return LEGACY[name]


async def reuse(owner, script, classname, methods):
    for method in methods:
        original = getattr(legacy(script), classname)(method)
        if hasattr(original, "setUp"):
            original.setUp()
        result = getattr(original, method)()
        if inspect.isawaitable(result):
            await result
        provenance = f"testing_modules/old_test/{script}::{classname}.{method}"
        REUSED.append(provenance)
        owner.note("historical assertion reuse", "unchanged method passes", provenance)


def case(test_id, objective, scenario, expected):
    def decorate(fn):
        @wraps(fn)
        async def bounded(self):
            return await asyncio.wait_for(fn(self),timeout=30)
        bounded.pre_uat = dict(test_id=test_id, objective=objective, scenario=scenario, expected=expected)
        return bounded
    return decorate


class EvidenceCase(unittest.IsolatedAsyncioTestCase):
    def setUp(self):
        self.observations = []

    def note(self, label, expected, actual):
        self.observations.append(dict(check=label, expected=expected, actual=actual))

    async def asyncSetUp(self):
        # Avoid asyncio slow-callback diagnostics for synchronous tiny workbook
        # creation. This affects test-loop diagnostics only, not worker settings.
        asyncio.get_running_loop().set_debug(False)


def oracle(raw):
    """Independent decimal arithmetic: no production calculation helper/constants.

    Use decimal-exact inputs away from ambiguous binary half-ties; HALF_EVEN
    is the documented two-decimal rounding convention used in this suite.
    """
    def d(value):
        return Decimal(str(value)) if value is not None else Decimal(0)
    scores = [raw.prelim_score, raw.midterm_score, raw.final_score]
    maxima = [raw.prelim_max, raw.midterm_max, raw.final_max]
    denominator = sum(map(d, maxima))
    direct = sum(map(d, scores)) / denominator if denominator > 0 else None
    indirect = d(raw.indirect_rating) / Decimal(5) * 100 if raw.indirect_rating is not None else None
    comp = None if direct is None else direct * 100 if indirect is None else Decimal("0.70") * direct * 100 + Decimal("0.30") * indirect
    comp = comp.quantize(Decimal("0.01"), rounding=ROUND_HALF_EVEN) if comp is not None else None
    level = None if comp is None else "Exceptional" if comp >= 85 else "Proficient" if comp >= 70 else "Basic" if comp >= 60 else "Below Basic"
    return dict(direct=None if direct is None else float(direct), indirect=None if indirect is None else float(indirect), composite=None if comp is None else float(comp), met=None if comp is None else comp >= 70, level=level, complete=all(s is not None for s in scores))


class SyntheticWorkbook:
    def __init__(self, students=4, clos=3, scenario="mixed"):
        self.students, self.clos, self.scenario = students, clos, scenario

    def __enter__(self):
        self.temp = tempfile.TemporaryDirectory(prefix="obelisk_pre_uat_")
        self.path = Path(self.temp.name) / "synthetic.xlsx"
        cls = legacy("test_etl_v2.py").TestETLV2Pipeline
        cls()._create_synthetic_aunobe_workbook(self.path, clo_count=self.clos, num_students=self.students, tamper_excel_attainment=True)
        with self.edit() as wb:
            direct, indirect, setup = wb[C.SheetNames.DIRECT_CLO], wb[C.SheetNames.INDIRECT_CLO], wb[C.SheetNames.SETUP]
            for coord, value in {"A3":"Course Title", "C3":"Course Code", "A4":"Section", "C4":"Term / AY", "A5":"Faculty Name"}.items():
                setup[coord] = value
            setup["B3"] = "Synthetic Pre-UAT Course"
            setup["D3"] = "SYN-101"
            setup["B5"] = "Synthetic Instructor"
            for s in range(self.students):
                dr, ir = C.DirectCloSheet.DATA_START_ROW+s, C.IndirectCloSheet.DATA_START_ROW+s
                direct.cell(dr, 1, f"SYN-{s+1:04d}")
                indirect.cell(ir, 1, f"SYN-{s+1:04d}")
                direct.cell(dr, 2, f"Synthetic Person {s+1:03d}")
                indirect.cell(ir, 2, f"Synthetic Person {s+1:03d}")
                for c in range(self.clos):
                    dc = C.DirectCloSheet.FIRST_BLOCK_START_COL + c*C.DirectCloSheet.BLOCK_WIDTH
                    ic = C.IndirectCloSheet.FIRST_BLOCK_START_COL + c*C.IndirectCloSheet.BLOCK_WIDTH
                    fraction, rating = self.values(s, c)
                    for offset, value in enumerate([fraction*25,25,fraction*50,50,fraction*25,25]):
                        direct.cell(dr, dc+offset, value)
                    indirect.cell(ir, ic, rating)
        return self

    def values(self, student, clo):
        if self.scenario == "worst": return 0.0, 1.0
        if self.scenario == "ideal": return 0.95, 5.0
        if self.scenario == "zero_gap": return 0.70, 3.5
        if self.scenario == "uniform": return 0.80, 4.5
        return [(0.40,2.0),(0.75,3.5),(0.90,4.5),(0.55,4.0)][(student+clo)%4]

    def edit(self):
        from contextlib import contextmanager
        @contextmanager
        def editor():
            wb = load_workbook(self.path)
            try:
                yield wb
                wb.save(self.path)
            finally:
                wb.close()
        return editor()

    def __exit__(self, *args):
        self.temp.cleanup()

    async def etl(self):
        extracted = await ExcelExtractor().extract(self.path)
        return extracted, await SimpleTransformer().transform(extracted)


def valid_recommendation(summary="Review the measured CLO gaps."):
    return {"summary":summary, "recommendations":[
        {"title":"Review aligned evidence", "explanation":"Review the assessment rubric with the faculty.", "recommendedFor":"Faculty"},
        {"title":"Discuss the next CQI cycle", "explanation":"Have a human review the evidence before deciding on action.", "recommendedFor":"Program Chair"},
    ], "pattern":None}


def validate_recommendation(value):
    """Independent evidence validator, NOT production response enforcement."""
    data = json.loads(value) if isinstance(value, str) else value
    if not isinstance(data, dict) or set(data) != {"summary","recommendations","pattern"}:
        raise ValueError("Invalid recommendation top-level schema")
    if not isinstance(data["summary"], str) or not data["summary"].strip():
        raise ValueError("Invalid summary")
    items = data["recommendations"]
    if not isinstance(items, list) or not 2 <= len(items) <= 3:
        raise ValueError("Expected two or three gap recommendations")
    for item in items:
        if not isinstance(item, dict) or set(item) != {"title","explanation","recommendedFor"}:
            raise ValueError("Invalid recommendation item")
        if any(not isinstance(item[k],str) or not item[k].strip() for k in item):
            raise ValueError("Invalid recommendation text type")
        if item["recommendedFor"] not in {"Faculty","Program Chair","Dean","VPAA","AQAU"}:
            raise ValueError("Invalid recommendedFor role")
    if data["pattern"] is not None and not isinstance(data["pattern"], str):
        raise ValueError("Invalid pattern type")
    return data


async def mocked_recommendation(header, records, text=None, mapping=None):
    payload = json.dumps(valid_recommendation()) if text is None else text
    with patch.object(CQI, "call_llm_api", new=AsyncMock(return_value=payload)) as mock:
        result = await CQI.generate_cqi_recommendation(header, records, mapping)
    return result, mock


def assert_oracle(owner, raw, output):
    expected = oracle(raw)
    owner.assertEqual((output.student_id, output.clo_code), (raw.student_id, raw.clo_code))
    for field, key in [("direct_clo_attainment_pct","direct"),("indirect_clo_attainment_pct","indirect"),("composite_clo_attainment_pct","composite")]:
        if expected[key] is None: owner.assertIsNone(getattr(output, field))
        else: owner.assertAlmostEqual(getattr(output, field),expected[key],places=8)
    owner.assertEqual(output.met_threshold, expected["met"])
    owner.assertEqual(output.clo_level, expected["level"])
    owner.assertEqual(output.is_record_complete, expected["complete"])
    owner.assertIsNone(output.excluded_reason)
    return expected


def gap_oracle(raw_records):
    groups = {}
    for raw in raw_records:
        group = groups.setdefault(raw.clo_code, {"total":0,"failed":[]})
        group["total"] += 1
        expected = oracle(raw)
        if expected["met"] is False: group["failed"].append(expected)
    return {clo:g for clo,g in groups.items() if g["failed"]}


CLAIM = re.compile(r"(CLO\d+): (\d+) of (\d+) students below (\d+(?:\.\d+)?)%; mean failed composite (\d+(?:\.\d+)?)%\.")


def check_numeric_claims(data, expected_gaps):
    """Check a deliberately narrow, documented grammar; flag all other digits.

    Does not pretend that arbitrary natural-language numbers are parseable.
    """
    checked, mismatches, manual = [], [], []
    texts = [data["summary"], data["pattern"] or ""]
    texts += [item[k] for item in data["recommendations"] for k in ("title","explanation")]
    for text in texts:
        for match in CLAIM.finditer(text):
            clo, failed, total, threshold, mean = match.groups()
            group = expected_gaps.get(clo)
            expected = None if group is None else [len(group["failed"]),group["total"],70.0, round(sum(x["composite"] for x in group["failed"])/len(group["failed"]),1)]
            actual = [int(failed),int(total),float(threshold),float(mean)]
            checked.append({"clo":clo,"expected":expected,"actual":actual})
            if actual != expected: mismatches.append(match.group())
        remainder = CLAIM.sub("",text)
        if re.search(r"\d", remainder): manual.append(remainder)
    return {"checked":checked,"mismatches":mismatches,"manual_review":manual,"semantic_review":"Human review required for educational usefulness, inferred causes, and advisory wording"}
