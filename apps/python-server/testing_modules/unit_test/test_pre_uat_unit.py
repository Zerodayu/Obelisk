"""UT-001..UT-010 and UT-014. Run only via the offline Phase 2A runner."""
import json
import tempfile
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import AsyncMock, Mock, patch

from testing_modules.pre_uat.support import (
    C, CQI, EvidenceCase, ExcelExtractor, SimpleTransformer, SyntheticWorkbook,
    assert_oracle, case, gap_oracle, mocked_recommendation, oracle, reuse,
    valid_recommendation, validate_recommendation,
)
from pydantic import ValidationError
from app.core.exceptions import OBELISKError, InvalidWorkbook, InvalidTemplate, MissingWorksheet, TransformationError, WorkbookValidationError
from app.etl.composite_attainment import compute_composite_clo


class UnitPreUAT(EvidenceCase):
    @case("UT-001", "Parse valid Excel class record", "Synthetic 4 students x 3 CLOs and historical dynamic 3/7-CLO cases", "Exact header, roster, CLO set, raw values and 12 independently checked outputs")
    async def test_ut001_valid_workbook(self):
        await reuse(self,"test_etl_v2.py","TestETLV2Pipeline",["test_dynamic_clo_count_discovery"])
        with SyntheticWorkbook() as fixture:
            extracted, output = await fixture.etl()
            header, raw, mapping = extracted
            self.assertEqual(header.course_code,"SYN-101")
            self.assertEqual(header.course_title,"Synthetic Pre-UAT Course")
            self.assertEqual(header.instructor_name,"Synthetic Instructor")
            self.assertEqual(header.no_of_students,4)
            self.assertEqual(extracted.section_extraction,"ok")
            self.assertEqual(extracted.section.code,"4A")
            self.assertEqual(extracted.section.program,"BSIT")
            self.assertEqual(extracted.setup.course_code,"SYN-101")
            self.assertEqual(mapping,[])
            self.assertEqual(len(raw),12)
            self.assertEqual({r.clo_code for r in raw},{"CLO1","CLO2","CLO3"})
            self.assertEqual({r.student_id for r in raw},{f"SYN-{i:04d}" for i in range(1,5)})
            for record in raw:
                student = int(record.student_id.removeprefix("SYN-"))-1
                clo = int(record.clo_code.removeprefix("CLO"))-1
                fraction,rating = fixture.values(student,clo)
                self.assertEqual(record.student_name,f"Synthetic Person {student+1:03d}")
                for actual,written in zip((record.prelim_score,record.prelim_max,record.midterm_score,record.midterm_max,record.final_score,record.final_max,record.indirect_rating),(fraction*25,25,fraction*50,50,fraction*25,25,rating),strict=True):
                    self.assertAlmostEqual(actual,written,places=8)
            checked = [assert_oracle(self,r,a) for r,a in zip(raw,output,strict=True)]
            self.note("parsed workbook",{"students":4,"clos":3,"records":12},{"students":header.no_of_students,"clos":len({r.clo_code for r in raw}),"records":len(output),"oracle_checked":len(checked)})

    @case("UT-002", "Reject corrupted/unsupported workbook", "Corrupt XLSX bytes, unsupported text file, synthetic old-layout workbook and invalid marker", "Structured InvalidWorkbook/MissingWorksheet/InvalidTemplate errors")
    async def test_ut002_invalid_workbook(self):
        with tempfile.TemporaryDirectory(prefix="obelisk_pre_uat_") as folder:
            for name in ["corrupt.xlsx","unsupported.txt"]:
                with self.subTest(file=name):
                    path = Path(folder)/name
                    path.write_bytes(b"synthetic data, not an Excel ZIP")
                    with self.assertRaises(InvalidWorkbook) as ctx:
                        await ExcelExtractor().extract(path)
                    data = ctx.exception.to_dict()
                    self.assertEqual(data["error_type"],"InvalidWorkbook")
                    self.assertTrue(data["details"]["reason"])
                    self.note(name,"InvalidWorkbook",data["error_type"])
        with SyntheticWorkbook() as fixture:
            with fixture.edit() as wb:
                wb[C.SheetNames.DIRECT_CLO][C.TemplateValidation.DIRECT_CLO_MARKER_CELL] = "WRONG MARKER"
            with self.assertRaises(InvalidTemplate) as ctx:
                await fixture.etl()
            self.assertEqual(ctx.exception.to_dict()["details"]["cell"],C.TemplateValidation.DIRECT_CLO_MARKER_CELL)
            self.note("wrong marker","InvalidTemplate",ctx.exception.to_dict())
        with SyntheticWorkbook() as fixture:
            with fixture.edit() as wb:
                wb[C.SheetNames.DIRECT_CLO].title = "OLD DATABASE"
            with self.assertRaises(MissingWorksheet) as ctx:
                await fixture.etl()
            self.assertEqual(ctx.exception.to_dict()["details"]["expected_sheet_name"],C.SheetNames.DIRECT_CLO)
            self.note("unsupported layout","MissingWorksheet",ctx.exception.to_dict())

    @case("UT-003", "Reject missing required workbook fields", "Delete each required sheet and each required marker value", "Structured errors for all four mandatory structural fields; SETUP metadata is optional by current contract")
    async def test_ut003_missing_required_fields(self):
        for sheet in C.TemplateValidation.REQUIRED_SHEETS:
            with self.subTest(missing_sheet=sheet), SyntheticWorkbook() as fixture:
                with fixture.edit() as wb: del wb[sheet]
                with self.assertRaises(MissingWorksheet) as ctx: await fixture.etl()
                self.assertEqual(ctx.exception.to_dict()["details"]["expected_sheet_name"],sheet)
                self.note("required sheet "+sheet,"MissingWorksheet",ctx.exception.to_dict())
        for sheet,marker in [(C.SheetNames.DIRECT_CLO,C.TemplateValidation.DIRECT_CLO_MARKER_CELL),(C.SheetNames.INDIRECT_CLO,C.TemplateValidation.INDIRECT_CLO_MARKER_CELL)]:
            with self.subTest(missing_marker=sheet), SyntheticWorkbook() as fixture:
                with fixture.edit() as wb: wb[sheet][marker] = None
                with self.assertRaises(InvalidTemplate) as ctx: await fixture.etl()
                self.assertEqual(ctx.exception.to_dict()["details"]["sheet_name"],sheet)
                self.note("required marker "+sheet,"InvalidTemplate",ctx.exception.to_dict())
        with SyntheticWorkbook() as fixture:
            with fixture.edit() as wb: del wb[C.SheetNames.SETUP]
            extracted, _ = await fixture.etl()
            self.assertEqual(extracted.section_extraction,"missing_sheet")
            self.assertIsNone(extracted[0].course_code)
            self.note("optional SETUP metadata","missing_sheet plus null metadata",{"section_extraction":extracted.section_extraction,"course_code":extracted[0].course_code})

    @case("UT-004", "Validate scores/IDs/CLO values while preserving extra credit", "Malformed numeric cells, duplicate identities/CLOs and ambiguous fallback; valid zero/blanks/numeric strings and extra credit", "Reject invalid nonblank cells and collisions with private structured errors; accept extra credit and preserve independently calculated outputs; curriculum validity is EXTERNAL")
    async def test_ut004_invalid_values(self):
        await reuse(self,"test_composite_attainment.py","TestCompositeAttainmentPureMath",["test_direct_scale_validation","test_indirect_scale_validation"])
        with SyntheticWorkbook(students=1,clos=1,scenario="uniform") as fixture:
            extracted, _ = await fixture.etl()
            header, records, mapping = extracted
            for field in ("prelim_score", "midterm_score", "final_score"):
                with self.subTest(regression="negative_"+field):
                    raw = records[0].model_copy(update={field:-0.01})
                    with self.assertRaises(TransformationError) as ctx:
                        await SimpleTransformer().transform((header,[raw],mapping))
                    details = ctx.exception.to_dict()
                    self.assertEqual(details["details"],{"field":field,"value":-0.01})
                    self.assertNotIn(raw.student_name,json.dumps(details))
                    self.assertNotIn(raw.student_id,json.dumps(details))
                    self.note("negative "+field,"private structured TransformationError",details)
                for value in (0.0,None,getattr(records[0],field)+100):
                    with self.subTest(compatibility=field,value=value):
                        raw = records[0].model_copy(update={field:value})
                        actual = (await SimpleTransformer().transform((header,[raw],mapping)))[0]
                        expected = assert_oracle(self,raw,actual)
                        self.note("zero/null/extra credit "+field,expected,{"composite":actual.composite_clo_attainment_pct,"complete":actual.is_record_complete})
        scenarios = ["negative_period_score","nonnumeric_score","nonnumeric_rating","duplicate_student_id","duplicate_clo_label"]
        for scenario in scenarios:
            with self.subTest(scenario=scenario), SyntheticWorkbook(students=2,clos=2,scenario="uniform") as fixture:
                with fixture.edit() as wb:
                    d,i = wb[C.SheetNames.DIRECT_CLO],wb[C.SheetNames.INDIRECT_CLO]
                    row,col = C.DirectCloSheet.DATA_START_ROW,C.DirectCloSheet.FIRST_BLOCK_START_COL
                    if scenario == "negative_period_score": d.cell(row,col,-1)
                    elif scenario == "nonnumeric_score": d.cell(row,col,"invalid-score")
                    elif scenario == "nonnumeric_rating": i.cell(C.IndirectCloSheet.DATA_START_ROW,C.IndirectCloSheet.FIRST_BLOCK_START_COL,"invalid-rating")
                    elif scenario == "duplicate_student_id":
                        d.cell(row+1,1,d.cell(row,1).value)
                        i.cell(C.IndirectCloSheet.DATA_START_ROW+1,1,i.cell(C.IndirectCloSheet.DATA_START_ROW,1).value)
                    else: d.cell(C.DirectCloSheet.CLO_LABEL_ROW,col+C.DirectCloSheet.BLOCK_WIDTH,"CLO1")
                error_type = TransformationError if scenario == "negative_period_score" else WorkbookValidationError
                with self.assertRaises(error_type) as ctx:
                    await fixture.etl()
                self.note(scenario,"structured rejection",ctx.exception.to_dict())
        await self._ut004_numeric_regressions()
        await self._ut004_identity_regressions()
        await self._ut004_clo_regressions()
        for field,offset,maximum in [("prelim_score",0,25),("midterm_score",2,50),("final_score",4,25)]:
            for bonus in (1,100):
                with self.subTest(extra_credit=field,bonus=bonus), SyntheticWorkbook(students=1,clos=1,scenario="uniform") as fixture:
                    with fixture.edit() as wb:
                        wb[C.SheetNames.DIRECT_CLO].cell(C.DirectCloSheet.DATA_START_ROW,C.DirectCloSheet.FIRST_BLOCK_START_COL+offset,maximum+bonus)
                    extracted,outputs = await fixture.etl()
                    expected = assert_oracle(self,extracted[1][0],outputs[0])
                    expected_direct = (80 - maximum*0.8 + maximum + bonus)/100
                    self.assertEqual(outputs[0].direct_clo_attainment_pct,expected_direct)
                    if bonus == 100: self.assertGreater(outputs[0].direct_clo_attainment_pct,1.0)
                    if field == "prelim_score" and bonus == 1:
                        self.assertEqual(outputs[0].direct_clo_attainment_pct,0.86)
                        self.assertEqual(outputs[0].composite_clo_attainment_pct,87.2)
                    self.note("extra credit accepted "+field,expected,{"direct":outputs[0].direct_clo_attainment_pct,"composite":outputs[0].composite_clo_attainment_pct})

    def _assert_workbook_error(self, error, sheet, cell, field):
        data = error.to_dict()
        self.assertEqual(data["error_type"],"WorkbookValidationError")
        self.assertEqual(data["details"]["sheet_name"],sheet)
        self.assertEqual(data["details"]["cell"],cell)
        self.assertEqual(data["details"]["field"],field)
        self.assertTrue(data["details"]["reason"])
        self.assertLessEqual(set(data["details"]),{"sheet_name","cell","field","reason","conflicting_cell"})
        for private in ("Synthetic Person","SYN-000","invalid-score","invalid-rating","private malformed text"):
            self.assertNotIn(private,json.dumps(data))
        return data

    async def _ut004_numeric_regressions(self):
        fields = [(C.SheetNames.DIRECT_CLO,0,"prelim_score"),(C.SheetNames.DIRECT_CLO,2,"midterm_score"),(C.SheetNames.DIRECT_CLO,4,"final_score"),(C.SheetNames.INDIRECT_CLO,0,"indirect_rating")]
        for sheet,offset,field in fields:
            row = C.DirectCloSheet.DATA_START_ROW if sheet == C.SheetNames.DIRECT_CLO else C.IndirectCloSheet.DATA_START_ROW
            column = (C.DirectCloSheet.FIRST_BLOCK_START_COL if sheet == C.SheetNames.DIRECT_CLO else C.IndirectCloSheet.FIRST_BLOCK_START_COL)+offset
            for value in ("private malformed text","#DIV/0!","#N/A",True,False,"NaN","Infinity","-inf"):
                with self.subTest(numeric_invalid=field,value_type=type(value).__name__,scenario=str(value)), SyntheticWorkbook(students=1,clos=1,scenario="uniform") as fixture:
                    with fixture.edit() as wb:
                        cell = wb[sheet].cell(row,column,value)
                        coordinate = cell.coordinate
                    with self.assertRaises(WorkbookValidationError) as ctx: await fixture.etl()
                    data = self._assert_workbook_error(ctx.exception,sheet,coordinate,field)
                    self.note("invalid numeric "+field,"private structured rejection",data)
            for value in (0,"0"," 4.5 ",None,""," \t "):
                with self.subTest(numeric_valid=field,scenario=str(value)), SyntheticWorkbook(students=1,clos=1,scenario="uniform") as fixture:
                    with fixture.edit() as wb: wb[sheet].cell(row,column).value = value
                    extracted,outputs = await fixture.etl()
                    parsed = getattr(extracted[1][0],field)
                    expected_value = None if value is None or isinstance(value,str) and not value.strip() else float(value)
                    self.assertEqual(parsed,expected_value)
                    expected = assert_oracle(self,extracted[1][0],outputs[0])
                    self.note("valid numeric/blank "+field,{"parsed":expected_value,**expected},{"parsed":parsed,"composite":outputs[0].composite_clo_attainment_pct})
        # Excel serialization replaces literal nonfinite floats with empty cells;
        # exercise these actual cell values before saving, and strings above through ETL.
        from openpyxl import Workbook
        wb = Workbook()
        sheet = wb.active
        sheet.title = C.SheetNames.DIRECT_CLO
        for value in (float("nan"),float("inf"),float("-inf")):
            with self.subTest(literal_nonfinite=str(value)):
                sheet.cell(5,3,value)
                with self.assertRaises(WorkbookValidationError) as ctx:
                    ExcelExtractor._read_numeric_cell(sheet,5,3,"prelim_score")
                self.note("literal nonfinite before serialization","rejected",self._assert_workbook_error(ctx.exception,sheet.title,"C5","prelim_score"))
        wb.close()
        self.assertIsNone(ExcelExtractor._as_optional_float("metadata remains permissive"))
        self.note("unrelated metadata parsing","unchanged permissive helper",None)

    async def _ut004_identity_regressions(self):
        for sheet in (C.SheetNames.DIRECT_CLO,C.SheetNames.INDIRECT_CLO):
            for same_name in (False,True):
                with self.subTest(duplicate_id=sheet,same_name=same_name), SyntheticWorkbook(students=2,clos=1) as fixture:
                    with fixture.edit() as wb:
                        ws = wb[sheet]
                        row = C.DirectCloSheet.DATA_START_ROW
                        ws.cell(row+1,1," "+ws.cell(row,1).value+" ")
                        if same_name: ws.cell(row+1,2,ws.cell(row,2).value)
                    with self.assertRaises(WorkbookValidationError) as ctx: await fixture.etl()
                    data = self._assert_workbook_error(ctx.exception,sheet,f"A{row+1}","student_id")
                    self.assertEqual(data["details"]["conflicting_cell"],f"A{row}")
                    self.note("duplicate ID "+sheet,"rejected before lookup overwrite",data)
        for mode in ("normal_ids","shared_name_ids","missing_ids","direct_missing_ids","unmatched_ids_name_fallback"):
            with self.subTest(valid_identity=mode), SyntheticWorkbook(students=2,clos=2) as fixture:
                with fixture.edit() as wb:
                    for s in range(2):
                        row = C.DirectCloSheet.DATA_START_ROW+s
                        if mode == "shared_name_ids":
                            for title in (C.SheetNames.DIRECT_CLO,C.SheetNames.INDIRECT_CLO): wb[title].cell(row,2,"Shared synthetic name")
                        if mode in ("missing_ids","direct_missing_ids"): wb[C.SheetNames.DIRECT_CLO].cell(row,1).value = None
                        if mode == "missing_ids": wb[C.SheetNames.INDIRECT_CLO].cell(row,1).value = None
                        if mode == "unmatched_ids_name_fallback": wb[C.SheetNames.INDIRECT_CLO].cell(row,1,"OTHER-"+str(s))
                extracted,outputs = await fixture.etl()
                self.assertEqual(len(outputs),4)
                for s,raw in enumerate(extracted[1]):
                    student = int(raw.student_id[-1])-1 if raw.student_id else int(raw.student_name.split()[-1])-1
                    clo = int(raw.clo_code[-1])-1
                    fraction,rating = fixture.values(student,clo)
                    self.assertAlmostEqual(raw.prelim_score,fraction*25)
                    self.assertEqual(raw.indirect_rating,rating)
                    assert_oracle(self,raw,outputs[s])
                self.note("identity compatibility "+mode,"four accurate distinct records",{"records":4,"ratings":[r.indirect_rating for r in extracted[1]]})
        for mode in ("ambiguous_indirect_name","duplicate_missing_id_name","ambiguous_direct_fallback"):
            with self.subTest(ambiguous_identity=mode), SyntheticWorkbook(students=2,clos=1) as fixture:
                with fixture.edit() as wb:
                    row = C.DirectCloSheet.DATA_START_ROW
                    if mode == "ambiguous_indirect_name":
                        wb[C.SheetNames.DIRECT_CLO].cell(row,1).value = None
                        wb[C.SheetNames.INDIRECT_CLO].cell(row+1,2,wb[C.SheetNames.INDIRECT_CLO].cell(row,2).value)
                    else:
                        for title in (C.SheetNames.DIRECT_CLO,C.SheetNames.INDIRECT_CLO):
                            wb[title].cell(row+1,2,wb[title].cell(row,2).value)
                        wb[C.SheetNames.DIRECT_CLO].cell(row,1).value = None
                        if mode == "duplicate_missing_id_name": wb[C.SheetNames.DIRECT_CLO].cell(row+1,1).value = None
                with self.assertRaises(WorkbookValidationError) as ctx: await fixture.etl()
                self.note(mode,"reject ambiguous fallback",self._assert_workbook_error(ctx.exception,C.SheetNames.DIRECT_CLO,f"B{row}","student_name"))

    async def _ut004_clo_regressions(self):
        for title in (C.SheetNames.DIRECT_CLO,C.SheetNames.INDIRECT_CLO):
            for whitespace in (False,True):
                with self.subTest(duplicate_clo=title,whitespace=whitespace), SyntheticWorkbook(students=2,clos=2) as fixture:
                    with fixture.edit() as wb:
                        ws = wb[title]
                        row = C.DirectCloSheet.CLO_LABEL_ROW if title == C.SheetNames.DIRECT_CLO else C.IndirectCloSheet.HEADER_ROW
                        col = C.DirectCloSheet.FIRST_BLOCK_START_COL if title == C.SheetNames.DIRECT_CLO else C.IndirectCloSheet.FIRST_BLOCK_START_COL
                        width = C.DirectCloSheet.BLOCK_WIDTH if title == C.SheetNames.DIRECT_CLO else C.IndirectCloSheet.BLOCK_WIDTH
                        label = "CLO1" if title == C.SheetNames.DIRECT_CLO else "CLO1 Different rating header"
                        ws.cell(row,col+width,"  "+label+"  " if whitespace else label)
                        coordinate,first = ws.cell(row,col+width).coordinate,ws.cell(row,col).coordinate
                    with self.assertRaises(WorkbookValidationError) as ctx: await fixture.etl()
                    data = self._assert_workbook_error(ctx.exception,title,coordinate,"clo_code")
                    self.assertEqual(data["details"]["conflicting_cell"],first)
                    self.note("duplicate CLO "+title,"rejected before distinct block values overwrite",data)

    @case("UT-005", "Independently verify 70/30 composite calculation", "Historical exact math/fallback/zero tests plus synthetic 80% direct, 90% indirect workbook", "83.00 composite for every workbook record; missing/zero indirect distinguished")
    async def test_ut005_composite(self):
        await reuse(self,"test_composite_attainment.py","TestCompositeAttainmentPureMath",["test_70_30_composite_calculation","test_fallback_when_indirect_is_none","test_indirect_zero_is_not_treated_as_missing"])
        await reuse(self,"test_etl_v2.py","TestIndirectAttainment",["test_compute_indirect_clo_attainment_returns_expected_percentage","test_compute_indirect_clo_attainment_raises_for_zero_scale_max","test_compute_indirect_clo_attainment_raises_for_exceeding_mean_rating"])
        await reuse(self,"test_composite_attainment.py","TestAunObeRealWorkbookIntegration",["test_no_indirect_sheet_or_scores_fallback_in_transformer"])
        with SyntheticWorkbook(scenario="uniform") as fixture:
            extracted,output = await fixture.etl()
            for raw, att in zip(extracted[1],output,strict=True):
                assert_oracle(self,raw,att)
                self.assertEqual(att.composite_clo_attainment_pct,83.0)
            self.note("70/30 workbook",83.0,[a.composite_clo_attainment_pct for a in output])
        self.assertIsNone(compute_composite_clo(None,80.0))
        self.note("no direct evidence",None,compute_composite_clo(None,80.0))

    @case("UT-006", "Exactly 70% attainment boundary", "69.99%, 70.00%, 70.01%, 69.996% rounded, custom course threshold 90%", "70.00 is met and Proficient, below is Basic; fixed institutional threshold")
    async def test_ut006_threshold(self):
        await reuse(self,"test_composite_attainment.py","TestCompositeAttainmentPureMath",["test_rounding_boundary","test_clo_levels_classification"])
        with SyntheticWorkbook(students=1,clos=1) as fixture:
            extracted,_ = await fixture.etl()
            header,raw,_ = extracted
            header = header.model_copy(update={"threshold":0.90})
            for pct,expected,met in [(69.99,69.99,False),(70.0,70.0,True),(70.01,70.01,True),(69.996,70.0,True)]:
                with self.subTest(direct_pct=pct):
                    record = raw[0].model_copy(update={"prelim_score":pct,"prelim_max":100.0,"midterm_score":None,"midterm_max":0.0,"final_score":None,"final_max":0.0,"indirect_rating":None})
                    att = (await SimpleTransformer().transform((header,[record],[])))[0]
                    self.assertEqual(att.composite_clo_attainment_pct,expected)
                    self.assertEqual(att.met_threshold,met)
                    self.assertEqual(att.clo_level,"Proficient" if met else "Basic")
                    self.note("threshold "+str(pct),{"composite":expected,"met":met},{"composite":att.composite_clo_attainment_pct,"met":att.met_threshold,"level":att.clo_level})

    @case("UT-007", "Numerical CLO data prepared for AI", "Mixed synthetic ETL records; anonymization and alignment in captured prompt", "All counts, composite/direct/indirect percentages, thresholds and means match independent oracle; no synthetic identities/IDs leak")
    async def test_ut007_ai_input(self):
        with SyntheticWorkbook() as fixture:
            extracted,output = await fixture.etl()
            before = [r.model_dump() for r in output]
            result,mock = await mocked_recommendation(extracted[0],output,mapping=[{"clo_code":"CLO1","blooms_level":"Apply","ipd_stage":"p","assessment_types":["Exam"]}])
            self.assertEqual(mock.await_count,1)
            prompt = mock.await_args.args[0]
            expected = gap_oracle(extracted[1])
            self.assertEqual(len(result["gaps"]),len(expected))
            for gap in result["gaps"]:
                group = expected[gap["clo_code"]]
                self.assertEqual(gap["total_students"],group["total"])
                self.assertEqual(gap["num_students_below_threshold"],len(group["failed"]))
                self.assertEqual(gap["attainment_values"],[e["composite"] for e in group["failed"]])
                self.assertEqual(gap["threshold"],70.0)
                for detail,exp in zip(gap["student_details"],group["failed"],strict=True):
                    self.assertEqual(detail["direct_pct"],round(exp["direct"]*100,2))
                    self.assertEqual(detail["indirect_pct"],exp["indirect"])
                    self.assertEqual(detail["composite_pct"],exp["composite"])
                mean = sum(e["composite"] for e in group["failed"])/len(group["failed"])
                self.assertIn(f"{len(group['failed'])} of {group['total']} students",prompt)
                self.assertIn(f"{mean:.1f}%",prompt)
            for raw in extracted[1]:
                self.assertNotIn(raw.student_name,prompt)
                self.assertNotIn(raw.student_id,prompt)
            anonymized = CQI.anonymize_students(output)
            self.assertTrue(all(r.student_id is None for r in anonymized))
            self.assertEqual([r.model_dump() for r in output],before)
            self.assertIn("Bloom's Apply",prompt)
            self.assertEqual(prompt.count("<DATA>"),1)
            self.assertEqual(prompt.count("</DATA>"),1)
            self.note("AI request",{"gaps":expected,"identity_leaks":0},{"gaps":result["gaps"],"prompt":prompt,"identity_leaks":0})

    @case("UT-008", "Validate AI recommendation response schema", "Valid JSON plus malformed JSON, missing summary, bad role/type, missing pattern, extra field", "Valid response preserved; invalid schemas rejected by service or status=error")
    async def test_ut008_response_schema(self):
        with SyntheticWorkbook() as fixture:
            extracted,output = await fixture.etl()
            result,_ = await mocked_recommendation(extracted[0],output)
            self.assertEqual(validate_recommendation(result["recommendation"]),valid_recommendation())
            for count,pattern in [(1,None),(2,"Measured cross-CLO pattern")]:
                payload = valid_recommendation()
                payload["recommendations"] = payload["recommendations"][:count]
                payload["pattern"] = pattern
                text = json.dumps(payload)
                for fenced in (False,True):
                    with self.subTest(compatibility="valid_json",items=count,fenced=fenced):
                        response,_ = await mocked_recommendation(extracted[0],output,"```json\n"+text+"\n```" if fenced else text)
                        self.assertEqual(response["status"],"ok")
                        self.assertIsInstance(response["recommendation"],str)
                        self.assertEqual(response["recommendation"],text)
                        self.assertEqual(json.loads(response["recommendation"]),payload)
                        self.note("valid fenced/single JSON","ok and original JSON string",{"items":count,"fenced":fenced,"status":response["status"]})
            bad = {"malformed_json":"not JSON","json_null":"null","json_array":"[]","json_scalar":"42","summary_type":json.dumps({**valid_recommendation(),"summary":False}),"recommendations_type":json.dumps({**valid_recommendation(),"recommendations":{}}),"item_type":json.dumps({**valid_recommendation(),"recommendations":[42]})}
            for scenario in ["missing_summary","bad_role","bad_type","missing_pattern","extra_field","bad_cardinality","bad_pattern","blank_title"]:
                payload = valid_recommendation()
                if scenario == "missing_summary": del payload["summary"]
                elif scenario == "bad_role": payload["recommendations"][0]["recommendedFor"] = "Automatic Approver"
                elif scenario == "bad_type": payload["recommendations"][0]["title"] = 42
                elif scenario == "missing_pattern": del payload["pattern"]
                elif scenario == "extra_field": payload["automaticApproval"] = True
                elif scenario == "bad_cardinality": payload["recommendations"] = []
                elif scenario == "bad_pattern": payload["pattern"] = 123
                else: payload["recommendations"][0]["title"] = "   "
                bad[scenario] = json.dumps(payload)
            for field in ("title","explanation","recommendedFor"):
                for mutation in ("missing","null","number","blank"):
                    payload = valid_recommendation()
                    if mutation == "missing": del payload["recommendations"][0][field]
                    else: payload["recommendations"][0][field] = {"null":None,"number":42,"blank":" "}[mutation]
                    bad[field+"_"+mutation] = json.dumps(payload)
            for scenario,text in bad.items():
                with self.subTest(scenario=scenario):
                    with self.assertRaises((ValueError,TypeError)): validate_recommendation(text)
                    try:
                        response,_ = await mocked_recommendation(extracted[0],output,text)
                    except (OBELISKError,ValidationError,ValueError):
                        self.note(scenario,"rejected","validation exception")
                    else:
                        self.note(scenario,"error status",{"status":response["status"],"recommendation":response["recommendation"]})
                        self.assertEqual(response["status"],"error","Invalid recommendation schema accepted: "+scenario)
                        self.assertEqual(response["error"],response["recommendation"])

    @case("UT-009", "Handle invalid/unavailable AI responses", "Mock SDK timeout, failover, rate-limit exhaustion, missing keys, blank response", "Failover order and error status correct; blank provider response must not be status=ok")
    async def test_ut009_ai_failures(self):
        settings = SimpleNamespace(llm_api_keys_list=["synthetic_slot_one","synthetic_slot_two"],LLM_MODEL="offline-test-model")
        generate = AsyncMock(side_effect=[TimeoutError("synthetic timeout"),SimpleNamespace(text=json.dumps(valid_recommendation()))])
        client = SimpleNamespace(aio=SimpleNamespace(models=SimpleNamespace(generate_content=generate)))
        with patch.object(CQI,"settings",settings),patch.object(CQI,"IS_DEBUG_MODE",False),patch.object(CQI.genai,"Client",return_value=client) as factory:
            response = await CQI.call_llm_api("Synthetic prompt")
            self.assertEqual(validate_recommendation(response),valid_recommendation())
            self.assertEqual(factory.call_count,2)
            self.assertEqual([c.kwargs["api_key"] for c in factory.call_args_list],settings.llm_api_keys_list)
            self.note("SDK failover",{"attempts":2,"status":"valid"},{"attempts":generate.await_count,"status":"valid","model":"offline-test-model"})
        with patch.object(CQI,"settings",SimpleNamespace(llm_api_keys_list=[])),patch.object(CQI,"IS_DEBUG_MODE",False):
            with self.assertRaises(NotImplementedError): await CQI.call_llm_api("Synthetic prompt")
            self.note("missing test keys","NotImplementedError","NotImplementedError")
        with SyntheticWorkbook() as fixture:
            extracted,output = await fixture.etl()
            generate = AsyncMock(side_effect=RuntimeError("429 synthetic rate limit"))
            client = SimpleNamespace(aio=SimpleNamespace(models=SimpleNamespace(generate_content=generate)))
            with patch.object(CQI,"settings",settings),patch.object(CQI,"IS_DEBUG_MODE",False),patch.object(CQI.genai,"Client",return_value=client):
                response = await CQI.generate_cqi_recommendation(extracted[0],output)
                self.assertEqual(response["status"],"error")
                self.assertEqual(generate.await_count,2)
                self.assertIn("error",response)
                self.note("all slots rate limited","error",{"status":response["status"],"attempts":generate.await_count})
            from app.analytics import institutional_summary as institutional
            from app.schemas.institutional_summary import InstitutionalSummaryPayload, CourseSubmission, Period
            payload = InstitutionalSummaryPayload(period=Period(type="semester",label="Synthetic"),submissions=[CourseSubmission(header=extracted[0],attainments=output,clo_plo_mapping=[])])
            baseline = institutional.compute_summary_only(payload.model_copy(deep=True))
            for label,text in [("empty",""),("whitespace"," \n\t"),("null",None),("number",42),("object",{}),("list",[]),("empty_fence","```json\n \n```"),("plain_fence","```\n\t```"),("text_fence","```text\n  \n```"),("bool",False)]:
                for boundary in ("course","institutional"):
                    with self.subTest(scenario=label,boundary=boundary):
                        if boundary == "course":
                            with patch.object(CQI,"call_llm_api",new=AsyncMock(return_value=text)) as mock:
                                response = await CQI.generate_cqi_recommendation(extracted[0],output)
                        else:
                            with patch.object(institutional,"call_llm_api",new=AsyncMock(return_value=text)) as mock:
                                response = await institutional.generate_institutional_summary(payload.model_copy(deep=True))
                            self.assertEqual(response["summary"],baseline)
                        self.assertEqual(response["status"],"error")
                        self.assertEqual(response["error"],response["recommendation"])
                        self.assertTrue(response["error"].startswith("[LLM API ERROR"))
                        self.assertEqual(mock.await_count,1)
                        self.note(boundary+" "+label,"error, one call, numerical summary preserved",{"status":response["status"],"error":response["error"],"calls":mock.await_count})
            for prose in ("Review measured program gaps with a human.","```\nReview measured program gaps with a human.\n```","[LLM API ERROR: Synthetic provider failure.]"):
                with patch.object(institutional,"call_llm_api",new=AsyncMock(return_value=prose)):
                    response = await institutional.generate_institutional_summary(payload.model_copy(deep=True))
                self.assertEqual(response["summary"],baseline)
                self.assertEqual(response["status"],"error" if prose.startswith("[LLM API ERROR") else "ok")
                self.assertEqual(response["recommendation"],CQI.strip_code_fences(prose))
                self.note("institutional prose/error compatibility","preserved text/summary",{"status":response["status"],"recommendation":response["recommendation"]})

    @case("UT-010", "Independent full-record calculation oracle", "Mixed synthetic 7 students x 4 CLOs, spoofed formula columns, missing scores and zero maxima", "All 28 records match independent direct/indirect/composite/threshold/level/completeness; section completeness correct")
    async def test_ut010_independent_oracle(self):
        await reuse(self,"test_composite_attainment.py","TestCompositeAttainmentPureMath",["test_mixed_composite_and_direct_only_mean_aggregation"])
        await reuse(self,"test_etl_v2.py","TestETLV2Pipeline",["test_independent_recomputation_ignores_tampered_formula_column"])
        with SyntheticWorkbook(students=7,clos=4) as fixture:
            with fixture.edit() as wb:
                d = wb[C.SheetNames.DIRECT_CLO]
                row,col = C.DirectCloSheet.DATA_START_ROW,C.DirectCloSheet.FIRST_BLOCK_START_COL
                d.cell(row,col).value = None
                for offset in [1,3,5]: d.cell(row,col+C.DirectCloSheet.BLOCK_WIDTH+offset,0)
            extracted,output = await fixture.etl()
            self.assertEqual(len(output),28)
            expectations = [assert_oracle(self,r,a) for r,a in zip(extracted[1],output,strict=True)]
            for att in output:
                group = [r for r in extracted[1] if r.clo_code==att.clo_code]
                complete = sum(oracle(r)["complete"] for r in group)/len(group)
                self.assertAlmostEqual(att.section_completeness_pct,complete)
                self.assertEqual(att.rule1_met,complete>=0.60)
            self.note("every record independently checked",28,{"checked":len(expectations),"expected":expectations,"actual":[a.model_dump() for a in output]})

    @case("UT-014", "150 unique student-CLO records", "Designated synthetic 30 students x 5 CLOs", "Exactly 30 IDs, five CLOs and 150 unique Cartesian pairs in extraction and transform; no duplicates")
    async def test_ut014_uniqueness(self):
        from collections import Counter
        with SyntheticWorkbook(students=30,clos=5) as fixture:
            extracted,output = await fixture.etl()
            expected = {(f"SYN-{s:04d}",f"CLO{c}") for s in range(1,31) for c in range(1,6)}
            for label,records in [("raw",extracted[1]),("transformed",output)]:
                counts = Counter((r.student_id,r.clo_code) for r in records)
                self.assertEqual(len(records),150)
                self.assertEqual(set(counts),expected)
                self.assertEqual(set(counts.values()),{1})
                self.assertEqual(len({r.student_id for r in records}),30)
                self.assertEqual(len({r.clo_code for r in records}),5)
                self.note(label,{"records":150,"unique_pairs":150,"duplicates":0},{"records":len(records),"unique_pairs":len(counts),"duplicates":sum(n-1 for n in counts.values())})
            for raw,att in zip(extracted[1],output,strict=True): assert_oracle(self,raw,att)
