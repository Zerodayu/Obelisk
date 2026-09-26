import tempfile
import unittest
from pathlib import Path
from openpyxl import Workbook

from app.etl import etl_const
from app.etl.indirect_attainment import compute_indirect_clo_attainment
from app.etl.extract.extractor import ExcelExtractor
from app.etl.transform.transformer import SimpleTransformer
from app.core.exceptions import MissingWorksheet


class TestIndirectAttainment(unittest.TestCase):
    """Unit tests for indirect attainment math helper."""

    def test_compute_indirect_clo_attainment_returns_expected_percentage(self):
        self.assertEqual(compute_indirect_clo_attainment(3.2, 4), 80.0)
        self.assertEqual(compute_indirect_clo_attainment(4.0, 5), 80.0)
        self.assertEqual(compute_indirect_clo_attainment(0, 5), 0.0)
        self.assertEqual(compute_indirect_clo_attainment(5, 5), 100.0)

    def test_compute_indirect_clo_attainment_raises_for_zero_scale_max(self):
        with self.assertRaises(ValueError):
            compute_indirect_clo_attainment(3.2, 0)

    def test_compute_indirect_clo_attainment_raises_for_exceeding_mean_rating(self):
        with self.assertRaises(ValueError):
            compute_indirect_clo_attainment(6, 5)


class TestETLV2Pipeline(unittest.IsolatedAsyncioTestCase):
    """
    Core test suite for ETL pipeline v2 (AUN-OBE template):
    1. Dynamic CLO count detection (synthetic file with 3 CLOs, another with 7 CLOs).
    2. Independent recalculation of Direct and Indirect CLO attainment from raw cells.
    3. Formula column spoofing resilience (modifying the Excel formula column does not alter calculated attainment).
    4. Old template format raises structured MissingWorksheet / InvalidTemplate error.
    """

    def setUp(self):
        self.extractor = ExcelExtractor()
        self.transformer = SimpleTransformer()

    def _create_synthetic_aunobe_workbook(
        self,
        dest_path: Path,
        clo_count: int = 3,
        num_students: int = 4,
        tamper_excel_attainment: bool = False,
    ):
        wb = Workbook()

        # Direct CLO sheet
        ws_d = wb.active
        ws_d.title = etl_const.SheetNames.DIRECT_CLO
        ws_d[etl_const.TemplateValidation.DIRECT_CLO_MARKER_CELL] = etl_const.TemplateValidation.DIRECT_CLO_MARKER_VALUE

        # Headers
        ws_d.cell(row=etl_const.DirectCloSheet.SUBFIELD_HEADER_ROW, column=1, value="Student ID")
        ws_d.cell(row=etl_const.DirectCloSheet.SUBFIELD_HEADER_ROW, column=2, value="Student Name")

        for c_idx in range(clo_count):
            clo_name = f"CLO{c_idx + 1}"
            start_col = etl_const.DirectCloSheet.FIRST_BLOCK_START_COL + (c_idx * etl_const.DirectCloSheet.BLOCK_WIDTH)
            ws_d.cell(row=etl_const.DirectCloSheet.CLO_LABEL_ROW, column=start_col, value=clo_name)
            ws_d.cell(row=etl_const.DirectCloSheet.SUBFIELD_HEADER_ROW, column=start_col + 0, value="Prelim Score")
            ws_d.cell(row=etl_const.DirectCloSheet.SUBFIELD_HEADER_ROW, column=start_col + 1, value="Prelim Max")
            ws_d.cell(row=etl_const.DirectCloSheet.SUBFIELD_HEADER_ROW, column=start_col + 2, value="Midterm Score")
            ws_d.cell(row=etl_const.DirectCloSheet.SUBFIELD_HEADER_ROW, column=start_col + 3, value="Midterm Max")
            ws_d.cell(row=etl_const.DirectCloSheet.SUBFIELD_HEADER_ROW, column=start_col + 4, value="Final Score")
            ws_d.cell(row=etl_const.DirectCloSheet.SUBFIELD_HEADER_ROW, column=start_col + 5, value="Final Max")
            ws_d.cell(row=etl_const.DirectCloSheet.SUBFIELD_HEADER_ROW, column=start_col + 6, value=f"{clo_name} Attainment %")

        # Indirect CLO sheet
        ws_i = wb.create_sheet(etl_const.SheetNames.INDIRECT_CLO)
        ws_i[etl_const.TemplateValidation.INDIRECT_CLO_MARKER_CELL] = etl_const.TemplateValidation.INDIRECT_CLO_MARKER_VALUE

        ws_i.cell(row=etl_const.IndirectCloSheet.HEADER_ROW, column=1, value="Student ID")
        ws_i.cell(row=etl_const.IndirectCloSheet.HEADER_ROW, column=2, value="Student Name")

        for c_idx in range(clo_count):
            clo_name = f"CLO{c_idx + 1}"
            start_col = etl_const.IndirectCloSheet.FIRST_BLOCK_START_COL + (c_idx * etl_const.IndirectCloSheet.BLOCK_WIDTH)
            ws_i.cell(row=etl_const.IndirectCloSheet.HEADER_ROW, column=start_col + 0, value=f"{clo_name} Rating (1-5)")
            ws_i.cell(row=etl_const.IndirectCloSheet.HEADER_ROW, column=start_col + 1, value=f"{clo_name} Attainment %")

        # Student data rows
        for s_idx in range(num_students):
            row = etl_const.DirectCloSheet.DATA_START_ROW + s_idx
            s_id = f"2026-{s_idx + 1:04d}"
            s_name = f"Student {chr(65 + s_idx)}"

            ws_d.cell(row=row, column=1, value=s_id)
            ws_d.cell(row=row, column=2, value=s_name)

            ws_i.cell(row=row, column=1, value=s_id)
            ws_i.cell(row=row, column=2, value=s_name)

            for c_idx in range(clo_count):
                d_start = etl_const.DirectCloSheet.FIRST_BLOCK_START_COL + (c_idx * etl_const.DirectCloSheet.BLOCK_WIDTH)
                # Raw scores: Prelim 20/25, Midterm 40/50, Final 20/25 -> Total 80/100 = 80% (0.80)
                ws_d.cell(row=row, column=d_start + 0, value=20.0)
                ws_d.cell(row=row, column=d_start + 1, value=25.0)
                ws_d.cell(row=row, column=d_start + 2, value=40.0)
                ws_d.cell(row=row, column=d_start + 3, value=50.0)
                ws_d.cell(row=row, column=d_start + 4, value=20.0)
                ws_d.cell(row=row, column=d_start + 5, value=25.0)

                # Attainment formula column: if tampered, put 999.0 to verify extractor ignores it
                fake_d_val = 999.0 if tamper_excel_attainment else 80.0
                ws_d.cell(row=row, column=d_start + 6, value=fake_d_val)

                i_start = etl_const.IndirectCloSheet.FIRST_BLOCK_START_COL + (c_idx * etl_const.IndirectCloSheet.BLOCK_WIDTH)
                # Indirect rating: 4.0 out of 5 -> (4 / 5) * 100 = 80.0%
                ws_i.cell(row=row, column=i_start + 0, value=4.0)
                fake_i_val = -500.0 if tamper_excel_attainment else 80.0
                ws_i.cell(row=row, column=i_start + 1, value=fake_i_val)

        # Summary rows in Direct and Indirect
        last_d_row = etl_const.DirectCloSheet.DATA_START_ROW + num_students
        ws_d.cell(row=last_d_row, column=2, value=f"{etl_const.DirectCloSheet.SUMMARY_ROW_LABEL_PREFIX} (Direct)")

        last_i_row = etl_const.IndirectCloSheet.DATA_START_ROW + num_students
        ws_i.cell(row=last_i_row, column=1, value=etl_const.IndirectCloSheet.SUMMARY_ROW_LABELS[0])

        # SETUP sheet
        ws_setup = wb.create_sheet(etl_const.SheetNames.SETUP)
        ws_setup["B3"] = "Synthetic Subject"
        ws_setup["D3"] = "SYN-101"
        ws_setup["B4"] = "BSIT-4A"
        ws_setup["D4"] = "2nd Sem, AY 2026-2027"
        ws_setup["B5"] = "Dr. Synthetic"
        ws_setup["H9"] = 70

        wb.save(dest_path)
        wb.close()

    async def test_dynamic_clo_count_discovery(self):
        """Verify dynamic discovery with 3 CLOs and 7 CLOs (not hardcoded to 5)."""
        with tempfile.TemporaryDirectory() as tmpdir:
            # Test 3 CLOs
            p3 = Path(tmpdir) / "synth_3_clos.xlsx"
            self._create_synthetic_aunobe_workbook(p3, clo_count=3, num_students=5)
            header3, records3, mapping3 = await self.extractor.extract(p3)
            self.assertEqual(len(records3), 5 * 3)  # 15 records
            unique_clos_3 = {r.clo_code for r in records3}
            self.assertEqual(unique_clos_3, {"CLO1", "CLO2", "CLO3"})
            self.assertEqual(len(mapping3), 0)

            # Test 7 CLOs
            p7 = Path(tmpdir) / "synth_7_clos.xlsx"
            self._create_synthetic_aunobe_workbook(p7, clo_count=7, num_students=3)
            header7, records7, mapping7 = await self.extractor.extract(p7)
            self.assertEqual(len(records7), 3 * 7)  # 21 records
            unique_clos_7 = {r.clo_code for r in records7}
            self.assertEqual(unique_clos_7, {f"CLO{i}" for i in range(1, 8)})

    async def test_independent_recomputation_ignores_tampered_formula_column(self):
        """
        Verify that direct and indirect attainment are independently recomputed from raw cells,
        completely ignoring any value in the workbook's formula columns.
        """
        with tempfile.TemporaryDirectory() as tmpdir:
            tampered_path = Path(tmpdir) / "tampered.xlsx"
            self._create_synthetic_aunobe_workbook(tampered_path, clo_count=2, num_students=2, tamper_excel_attainment=True)

            extracted = await self.extractor.extract(tampered_path)
            attainments = await self.transformer.transform(extracted)

            self.assertEqual(len(attainments), 4)
            for att in attainments:
                # Direct attainment should be exactly 80.0 / 100.0 = 0.80, NOT 999.0
                self.assertAlmostEqual(att.direct_clo_attainment_pct, 0.80, places=4)
                # Indirect attainment should be exactly (4 / 5) * 100 = 80.0%, NOT -500.0
                self.assertAlmostEqual(att.indirect_clo_attainment_pct, 80.0, places=4)
                self.assertTrue(att.met_threshold)
                self.assertEqual(att.clo_level, "Proficient")
                self.assertIsNone(att.excluded_reason)

    async def test_old_template_format_raises_structured_missing_worksheet_error(self):
        """Verify that an old template file raises MissingWorksheet instead of partially extracting."""
        old_template_path = Path("classrecord_templates/E-classrecord(LECTURE ONLY).xlsx")
        if old_template_path.exists():
            with self.assertRaises(MissingWorksheet) as ctx:
                await self.extractor.extract(old_template_path)
            err_dict = ctx.exception.to_dict()
            self.assertEqual(err_dict["error_type"], "MissingWorksheet")
            self.assertEqual(err_dict["details"]["expected_sheet_name"], etl_const.SheetNames.DIRECT_CLO)


if __name__ == "__main__":
    unittest.main()
