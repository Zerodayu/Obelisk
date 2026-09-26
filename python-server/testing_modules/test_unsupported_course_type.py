import asyncio
import tempfile
import unittest
from pathlib import Path
from unittest.mock import AsyncMock, patch

from openpyxl import Workbook

from app.etl import etl_const
from app.workers import worker


class TestUnsupportedCourseType(unittest.IsolatedAsyncioTestCase):
    async def test_research_course_type_marks_job_failed_with_structured_error(self):
        with tempfile.TemporaryDirectory() as tmpdir:
            workbook_path = Path(tmpdir) / "unsupported_course_type.xlsx"
            wb = Workbook()

            # Direct CLO sheet with valid marker
            ws_d = wb.active
            ws_d.title = etl_const.SheetNames.DIRECT_CLO
            ws_d[etl_const.TemplateValidation.DIRECT_CLO_MARKER_CELL] = etl_const.TemplateValidation.DIRECT_CLO_MARKER_VALUE
            ws_d.cell(row=etl_const.DirectCloSheet.CLO_LABEL_ROW, column=etl_const.DirectCloSheet.FIRST_BLOCK_START_COL, value="CLO1")
            ws_d.cell(row=etl_const.DirectCloSheet.DATA_START_ROW, column=1, value="S001")
            ws_d.cell(row=etl_const.DirectCloSheet.DATA_START_ROW, column=2, value="Test Student")

            # Indirect CLO sheet with valid marker
            ws_i = wb.create_sheet(etl_const.SheetNames.INDIRECT_CLO)
            ws_i[etl_const.TemplateValidation.INDIRECT_CLO_MARKER_CELL] = etl_const.TemplateValidation.INDIRECT_CLO_MARKER_VALUE
            ws_i.cell(row=etl_const.IndirectCloSheet.HEADER_ROW, column=etl_const.IndirectCloSheet.FIRST_BLOCK_START_COL, value="CLO1 Rating (1-5)")

            # SETUP sheet with unsupported course type 'RESEARCH'
            ws_setup = wb.create_sheet(etl_const.SheetNames.SETUP)
            ws_setup["B3"] = "Synthetic Test Course"
            ws_setup["D3"] = "TEST-101"
            ws_setup["B4"] = "A"
            ws_setup["D4"] = "SY 2025-2026, 1st Sem"
            ws_setup["B5"] = "Test Instructor"
            ws_setup["B6"] = "RESEARCH"  # Course type

            wb.save(workbook_path)
            wb.close()

            update_calls = []

            async def fake_update_job(job_id, updates):
                update_calls.append((job_id, updates))

            with patch("app.workers.worker.redis_client", object()), \
                 patch("app.workers.worker.get_job_from_queue", new=AsyncMock(side_effect=["job-1", asyncio.CancelledError()])), \
                 patch("app.workers.worker.get_job", new=AsyncMock(return_value={"payload": {"file_path": str(workbook_path)}})), \
                 patch("app.workers.worker.update_job", new=AsyncMock(side_effect=fake_update_job)):
                await worker.start_worker(worker_id=1)

            failed_updates = [updates for _, updates in update_calls if updates.get("status") == "failed"]
            self.assertEqual(len(failed_updates), 1)

            error_payload = failed_updates[0]["error"]
            self.assertEqual(error_payload["error_type"], "UnsupportedCourseType")
            self.assertEqual(
                error_payload["message"],
                "Course type 'RESEARCH' is not yet supported. Only LECTURE course records can be processed at this time.",
            )
            self.assertEqual(error_payload["details"], {"course_type": "RESEARCH", "supported_types": ["LECTURE"]})
