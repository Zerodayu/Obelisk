import json
import time
from tempfile import NamedTemporaryFile
from pathlib import Path
import sys
import requests
from openpyxl import load_workbook

# Add project root to path to allow imports from `app`
PROJECT_ROOT = Path(__file__).parent.parent
sys.path.insert(0, str(PROJECT_ROOT))

from app.etl import etl_const

# --- Configuration ---
BASE_URL = "http://localhost:8000"
TEMPLATES_DIR = PROJECT_ROOT / "classrecord_templates"
TEST_FILE = TEMPLATES_DIR / "E-classrecord(LECTURE ONLY).xlsx"

POLL_INTERVAL_SECONDS = 1
TIMEOUT_SECONDS = 30


def print_step(message: str):
    """Prints a formatted step message."""
    print(f"\n--- {message} ---")


def print_json(data: dict, title: str = "JSON Response"):
    """Prints a dictionary as formatted JSON."""
    print(f"{title}:")
    print(json.dumps(data, indent=2))


def create_unmapped_workbook_copy(source_path: Path) -> Path:
    """Creates a temporary workbook copy with one CLO row's PLO correlations zeroed out."""
    workbook = load_workbook(source_path)
    coverpage = workbook[etl_const.SheetNames.COVERPAGE]

    for col in range(etl_const.CloPlo.PLO_START_COL, coverpage.max_column + 1):
        plo_header = coverpage.cell(row=etl_const.CloPlo.PLO_HEADER_ROW, column=col).value
        if plo_header is None or str(plo_header).strip() == "":
            break
        coverpage.cell(row=etl_const.CloPlo.FIRST_CLO_ROW, column=col).value = 0

    tmp = NamedTemporaryFile(delete=False, suffix=".xlsx")
    tmp.close()
    temp_path = Path(tmp.name)
    workbook.save(temp_path)
    workbook.close()
    return temp_path


def main():
    """Runs the end-to-end test for file upload and processing."""
    print_step(f"Starting E2E test: Uploading file to {BASE_URL}")

    if not TEST_FILE.exists():
        print(f"ERROR: Test file not found at '{TEST_FILE}'")
        print("\nEND-TO-END TEST: FAILED")
        return

    temp_file = create_unmapped_workbook_copy(TEST_FILE)
    print(f"Using temporary workbook with one zeroed CLO-PLO mapping row: {temp_file}")

    try:
        # 1. POST the file to /upload
        try:
            with open(temp_file, "rb") as f:
                files = {"file": (temp_file.name, f, "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet")}
                response = requests.post(f"{BASE_URL}/upload", files=files, timeout=10)
        except requests.exceptions.ConnectionError as e:
            print(f"Connection Error: Could not connect to the server at {BASE_URL}.")
            print("Please make sure the server is running: uvicorn app.main:app --reload")
            print(f"\nEND-TO-END TEST: FAILED")
            return

        print(f"POST /upload status code: {response.status_code}")
        job_data = response.json()
        print_json(job_data, title="Initial Job Response")

        if response.status_code != 202 or "job_id" not in job_data:
            print("ERROR: Did not receive a 202 status and a valid job_id.")
            print("\nEND-TO-END TEST: FAILED")
            return

        job_id = job_data["job_id"]

        # 2. Poll GET /jobs/{job_id} for the result
        print_step(f"Polling for job '{job_id}' completion...")
        start_time = time.time()
        final_job_status = None

        while time.time() - start_time < TIMEOUT_SECONDS:
            print(".", end="", flush=True)
            time.sleep(POLL_INTERVAL_SECONDS)

            try:
                poll_response = requests.get(f"{BASE_URL}/jobs/{job_id}", timeout=5)
            except requests.exceptions.ConnectionError:
                print("\nERROR: Connection to server lost during polling.")
                print("\nEND-TO-END TEST: FAILED")
                return

            if poll_response.status_code != 200:
                print(f"\nWARN: Polling returned status {poll_response.status_code}")
                continue

            current_job = poll_response.json()
            status = current_job.get("status")

            if status in ["completed", "failed"]:
                final_job_status = current_job
                break

        print() # Newline after polling dots

        if not final_job_status:
            print(f"ERROR: Job did not complete within the {TIMEOUT_SECONDS}s timeout.")
            print("\nEND-TO-END TEST: FAILED")
            return

        # 3. Check for failure
        if final_job_status.get("status") == "failed":
            print_step("Job Failed")
            print_json(final_job_status or {}, title="Failed Job Details")
            print("\nEND-TO-END TEST: FAILED")
            return

        # 4. Process and verify the successful result
        print_step("Job Completed Successfully")
        result = final_job_status.get("result", {}).get("loaded", {})

        if not result:
            print("ERROR: 'result.loaded' field is missing or empty.")
            print_json(final_job_status or {}, "Full Job Status")
            print("\nEND-TO-END TEST: FAILED")
            return

        print_json(result, title="Final Loaded Result")

        attainments = result.get("attainments", [])
        num_attainments = len(attainments)
        print(f"\nFound {num_attainments} attainment records.")

        if num_attainments == 0:
            print("ERROR: No attainment records were found in the result.")
            print("\nEND-TO-END TEST: FAILED")
            return

        excluded_rows = [record for record in attainments if record.get("excluded_reason") == "no_plo_mapping"]
        if not excluded_rows:
            print("ERROR: Expected at least one excluded CLO row with 'no_plo_mapping'.")
            print("\nEND-TO-END TEST: FAILED")
            return

        print(f"✓ Found {len(excluded_rows)} excluded CLO attainment row(s) with excluded_reason='no_plo_mapping'.")

        # 5. Call the new recommendation endpoint
        print_step(f"5. Calling GET /analytics/jobs/{job_id}/recommendation")
        recommendation_url = f"{BASE_URL}/analytics/jobs/{job_id}/recommendation"
        try:
            rec_response = requests.get(recommendation_url, timeout=20) # Increased timeout for LLM call
            print(f"GET {recommendation_url} status code: {rec_response.status_code}")

            if rec_response.status_code == 200:
                rec_data = rec_response.json()
                print_json(rec_data, title="CQI Recommendation Response")

                recommendation_text = rec_data.get("recommendation", "")
                if recommendation_text and "LLM API ERROR" not in recommendation_text:
                    print("\nVERIFICATION PASSED: A real, non-error recommendation was received from the LLM.")
                    print("\nEND-TO-END TEST: PASSED")
                else:
                    print("\nVERIFICATION FAILED: Recommendation was empty or contained an API error.")
                    print("\nEND-TO-END TEST: FAILED")

            else:
                print(f"ERROR: Received status {rec_response.status_code}")
                try:
                    print_json(rec_response.json(), title="Error Response")
                except json.JSONDecodeError:
                    print("Could not decode JSON from error response.")
                    print(rec_response.text)
                print("\nEND-TO-END TEST: FAILED")

        except requests.exceptions.RequestException as e:
            print(f"ERROR: Request to recommendation endpoint failed: {e}")
            print("\nEND-TO-END TEST: FAILED")
    finally:
        if temp_file.exists():
            temp_file.unlink()


if __name__ == "__main__":
    main()